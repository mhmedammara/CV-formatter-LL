"""Remplissage du modèle PowerPoint Logiclever à partir des données vérifiées."""

from __future__ import annotations

import copy
import dataclasses
import re
from collections.abc import Callable
from dataclasses import dataclass, field
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.presentation import Presentation as PresentationFile
from pptx.slide import Slide

from . import contenu, fonts
from .config import (
    BLACK,
    BLUE,
    CONTENT_BOTTOM,
    CONTINUATION_HEADER_Y,
    FULL_WIDTH,
    LEFT_TOP_PAGE1,
    LEFT_WIDTH,
    LEFT_X,
    NO_PHOTO_X,
    PHOTO_BOTTOM,
    PHOTO_GAP,
    PHOTO_SIZE,
    PHOTO_X,
    RIGHT_WIDTH,
    RIGHT_X,
    SECTION_HEADER_HEIGHT,
    TEXT_INSET,
    WHITE,
    cm,
    emu_to_cm,
)
from .layout import PT_PER_CM, Block, Para, Run, box_height, fits_one_line, line_count, para_height, scale, text_width
from .schema import CV, Certification, Experience, Formation, Langue

# Élément XML de lxml : sa classe publique porte un nom « privé » (_Element) dans les annotations de lxml.
Element = etree._Element  # pyright: ignore[reportPrivateUsage]

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}


def q(tag: str) -> str:
    prefix, name = tag.split(":")
    return f"{{{NS[prefix]}}}{name}"


SHAPE_TAGS = {q("p:sp"), q("p:grpSp"), q("p:pic"), q("p:cxnSp")}
CONTACT_FIELDS = ("email", "telephone", "localisation", "linkedin")
RIGHT_BLOCKS = (("expertise", "EXPERTISE"), ("langues", "LANGUES"), ("si", "SI & outils"))
# Une liste de mots-clés trop courte met en avant des détails (« Scrum · AMDEC ») : en dessous de ce nombre
# d'éléments, le bloc n'est pas affiché. Les langues, certifications et diplômes ne sont pas concernés.
MIN_KEYWORD_ITEMS = 3
KEYWORD_BLOCKS = {"expertise": "Expertise", "si": "SI & outils"}
HEADER_GAP = 0.25  # espace avant un titre de section placé dans le flux
CONTACT_RIGHT_EDGE = 20.75


# --- Primitives XML ------------------------------------------------------------------


def child(el: Element, path: str) -> Element:
    """Sous-élément obligatoire du modèle (erreur explicite s'il manque)."""
    found = el.find(path, NS)
    if found is None:
        raise ValueError(f"Modèle PowerPoint inattendu : élément {path} introuvable")
    return found


def _xfrm(el: Element) -> Element:
    return child(el, "p:grpSpPr/a:xfrm" if el.tag == q("p:grpSp") else "p:spPr/a:xfrm")


def _emu(el: Element, attr: str) -> int:
    return int(el.get(attr, "0"))


def get_geom(el: Element) -> tuple[float, float, float, float]:
    xfrm = _xfrm(el)
    off, ext = child(xfrm, "a:off"), child(xfrm, "a:ext")
    return emu_to_cm(_emu(off, "x")), emu_to_cm(_emu(off, "y")), emu_to_cm(_emu(ext, "cx")), emu_to_cm(_emu(ext, "cy"))


def set_geom(el: Element, x: float | None = None, y: float | None = None, w: float | None = None, h: float | None = None) -> None:
    xfrm = _xfrm(el)
    off, ext = child(xfrm, "a:off"), child(xfrm, "a:ext")
    if x is not None:
        off.set("x", str(cm(x)))
    if y is not None:
        off.set("y", str(cm(y)))
    if w is not None:
        ext.set("cx", str(cm(w)))
    if h is not None:
        ext.set("cy", str(cm(h)))


def remove(el: Element | None) -> None:
    parent = el.getparent() if el is not None else None
    if el is not None and parent is not None:
        parent.remove(el)


def text_of(el: Element) -> str:
    return "".join(t.text or "" for t in el.iter(q("a:t")))


def _rpr(run: Run, tag: str = "a:rPr") -> Element:
    rpr = etree.Element(q(tag))
    rpr.set("lang", "fr-FR")
    rpr.set("sz", str(int(round(run.size * 100))))
    rpr.set("b", "1" if run.style == "bold" else "0")
    rpr.set("i", "0")
    rpr.set("dirty", "0")
    fill = etree.SubElement(rpr, q("a:solidFill"))
    etree.SubElement(fill, q("a:srgbClr"), val=run.color)
    face = fonts.typeface(run.style)
    for t in ("a:latin", "a:ea", "a:cs", "a:sym"):
        etree.SubElement(rpr, q(t), typeface=face)
    return rpr


def _paragraph(para: Para) -> Element:
    p = etree.Element(q("a:p"))
    ppr = etree.SubElement(p, q("a:pPr"))
    margin = cm(para.indent)
    ppr.set("marL", str(margin))
    ppr.set("indent", str(-margin if para.bullet else 0))
    ppr.set("algn", para.align)
    ppr.set("rtl", "0")
    etree.SubElement(etree.SubElement(ppr, q("a:lnSpc")), q("a:spcPct"), val=str(int(round(para.line_spacing * 100000))))
    etree.SubElement(etree.SubElement(ppr, q("a:spcBef")), q("a:spcPts"), val=str(int(round(para.space_before * 100))))
    etree.SubElement(etree.SubElement(ppr, q("a:spcAft")), q("a:spcPts"), val=str(int(round(para.space_after * 100))))
    if para.bullet:
        etree.SubElement(etree.SubElement(ppr, q("a:buClr")), q("a:srgbClr"), val=para.bullet_color)
        etree.SubElement(ppr, q("a:buSzPct"), val="100000")
        etree.SubElement(ppr, q("a:buFont"), typeface="Arial")
        etree.SubElement(ppr, q("a:buChar"), char=para.bullet)
    else:
        etree.SubElement(ppr, q("a:buNone"))
    runs = [r for r in para.runs if r.text]
    for run in runs:
        r = etree.SubElement(p, q("a:r"))
        r.append(_rpr(run))
        etree.SubElement(r, q("a:t")).text = run.text
    p.append(_rpr(runs[-1] if runs else Run(""), "a:endParaRPr"))
    return p


def set_text(sp: Element, paras: list[Para], anchor: str = "t") -> None:
    body = child(sp, "p:txBody")
    body_pr = child(body, "a:bodyPr")
    for key in ("lIns", "tIns", "rIns", "bIns"):
        body_pr.set(key, str(cm(TEXT_INSET)))
    body_pr.set("anchor", anchor)
    body_pr.set("wrap", "square")
    for node in list(body_pr):
        if node.tag in (q("a:normAutofit"), q("a:spAutoFit"), q("a:noAutofit")):
            body_pr.remove(node)
    body_pr.insert(0, etree.Element(q("a:noAutofit")))
    for p in body.findall("a:p", NS):
        body.remove(p)
    paras = list(paras) or [Para([Run("")])]
    # PowerPoint applique l'espace avant même au premier paragraphe : on l'annule (mesure cohérente).
    paras[0] = dataclasses.replace(paras[0], space_before=0)
    for para in paras:
        body.append(_paragraph(para))


# --- Modèle -------------------------------------------------------------------------


def prepare_presentation(prs: PresentationFile) -> None:
    """Nettoie le modèle : mises en page inutilisées, sous-ensembles de polices intégrés."""
    for master in prs.slide_masters:
        for layout in list(master.slide_layouts):
            if not layout.used_by_slides:
                master.slide_layouts.remove(layout)
    pres = prs.element
    font_list = pres.find(q("p:embeddedFontLst"))
    if font_list is not None:
        rel_ids = {rel for el in font_list.iter() if (rel := el.get(q("r:id")))}
        pres.remove(font_list)
        for rel_id in rel_ids:
            prs.part.drop_rel(rel_id)
    # PowerPoint réintègre les polices complètes à l'enregistrement (cf. export_pdf).
    pres.set("embedTrueTypeFonts", "1")
    pres.set("saveSubsetFonts", "0")


class Template:
    def __init__(self, prs: PresentationFile):
        self.prs = prs
        self.slide: Slide = prs.slides[0]
        self.tree: Element = self.slide.shapes.element
        top: list[Element] = [el for el in self.tree if el.tag in SHAPE_TAGS]
        self.placeholders: dict[str, Element] = {}
        for el in top:
            for name in re.findall(r"\{\{\s*(\w+)\s*\}+", text_of(el)):
                self.placeholders.setdefault(name, el)
        self.headers: dict[str, Element] = {
            text_of(el).strip(): el for el in top if el.tag == q("p:grpSp") and text_of(el).strip()
        }
        self.photo: Element | None = next((el for el in top if el.tag == q("p:pic")), None)
        self.icons = self._contact_icons(top)
        self._next_id = 1000

    def _contact_icons(self, top: list[Element]) -> dict[str, Element]:
        """Associe à chaque coordonnée l'icône sans texte la plus proche sur la même ligne."""
        candidates = [el for el in top if not text_of(el).strip() and el.tag != q("p:pic")]
        icons: dict[str, Element] = {}
        for name in CONTACT_FIELDS:
            box = self.placeholders.get(name)
            if box is None:
                continue
            bx, by, _, bh = get_geom(box)
            center = by + bh / 2
            best: Element | None = None
            best_dist = 0.8
            for el in candidates:
                x, y, w, h = get_geom(el)
                if x + w > bx + 0.1 or bx - x > 1.6:
                    continue
                dist = abs(y + h / 2 - center)
                if dist < best_dist:
                    best, best_dist = el, dist
            if best is not None:
                icons[name] = best
                candidates.remove(best)
        return icons

    def new_id(self) -> int:
        self._next_id += 1
        return self._next_id

    def clone(self, el: Element, tree: Element) -> Element:
        copy_el = copy.deepcopy(el)
        for nv in copy_el.iter(q("p:cNvPr")):
            nv.set("id", str(self.new_id()))
        tree.append(copy_el)
        return copy_el


def _retitle_header(group: Element, title: str) -> None:
    text_sp = next(sp for sp in group.iter(q("p:sp")) if text_of(sp).strip())
    runs = list(text_sp.iter(q("a:t")))
    runs[0].text = title
    for extra in runs[1:]:
        extra.text = ""
    rpr = text_sp.find(".//a:rPr", NS)
    size = int(rpr.get("sz", "1800")) / 100 if rpr is not None else 18.0
    needed = text_width(title, "medium", size) / PT_PER_CM + 2 * TEXT_INSET + 0.2
    sp_ext = child(child(text_sp, "p:spPr/a:xfrm"), "a:ext")
    delta = needed - emu_to_cm(_emu(sp_ext, "cx"))
    if delta > 0:
        sp_ext.set("cx", str(_emu(sp_ext, "cx") + cm(delta)))
        g_xfrm = child(group, "p:grpSpPr/a:xfrm")
        for tag in ("a:ext", "a:chExt"):
            ext = child(g_xfrm, tag)
            ext.set("cx", str(_emu(ext, "cx") + cm(delta)))


# --- En-tête : photo, nom, titre, pastille, coordonnées ----------------------------


def _place_photo(t: Template, photo: Path | None) -> bool:
    if t.photo is None:
        return False
    blip = child(t.photo, ".//a:blip")
    old_rel = blip.get(q("r:embed"), "")
    if photo is not None:
        _, rel_id = t.slide.part.get_or_add_image_part(str(photo))
        blip.set(q("r:embed"), rel_id)
        if old_rel:
            t.slide.part.drop_rel(old_rel)
        child(t.photo, "p:nvPicPr/p:cNvPr").set("descr", "Photo du consultant")
        # Photo détachée du bord de la feuille, alignée sur la marge du texte et sur le bas de la pastille.
        set_geom(t.photo, x=PHOTO_X, y=PHOTO_BOTTOM - PHOTO_SIZE, w=PHOTO_SIZE, h=PHOTO_SIZE)
        return True
    remove(t.photo)
    if old_rel:
        t.slide.part.drop_rel(old_rel)
    return False


def display_name(cv: CV) -> str:
    def nice(part: str) -> str:
        if part.isupper() or part.islower():
            return "-".join(w[:1].upper() + w[1:].lower() for w in part.split("-"))
        return part

    first = " ".join(nice(p) for p in (cv.prenom or "").split())
    last = (cv.nom or "").strip().upper()
    return " ".join(x for x in (first, last) if x) or "Consultant"


def _short_title(title: str) -> str:
    for sep in (" — ", " – ", " | ", "|", " - ", " / "):
        if sep in title:
            return title.split(sep)[0].strip()
    return title


def _place_identity(t: Template, name: str, title: str | None, label: str | None, has_photo: bool, notes: list[str]) -> None:
    name_box = t.placeholders["nom"]
    pill = t.placeholders.get("experience_label")
    nx, _, nw, _ = get_geom(name_box)
    pill_y = get_geom(pill)[1] if pill is not None else 5.32
    right = nx + nw
    nx = (PHOTO_X + PHOTO_SIZE + PHOTO_GAP if has_photo else NO_PHOTO_X) - TEXT_INSET
    nw = right - nx
    top, bottom = 1.45, pill_y - 0.1
    available = bottom - top

    def build(name_size: float, title_size: float, title_text: str | None) -> list[Para]:
        paras = [Para([Run(name, "medium", name_size, BLACK)], line_spacing=0.9)]
        if title_text:
            paras.append(Para([Run(title_text, "medium", title_size, BLUE)], space_before=4, line_spacing=1.0))
        return paras

    def name_lines(size: float) -> int | None:
        """Nombre de lignes du nom à cette taille, ou None si un mot ne tient pas sur une ligne."""
        if not all(fits_one_line(w, "medium", size, nw) for w in name.split()):
            return None
        return line_count(Para([Run(name, "medium", size, BLACK)], line_spacing=0.9), nw)

    # Le nom sur une seule ligne est préféré (on réduit jusqu'à 28 pt) ; sinon deux, puis trois lignes.
    name_sizes = [s for s in range(37, 27, -1) if name_lines(s) == 1]
    name_sizes += [s for s in range(37, 21, -1) if name_lines(s) == 2 and s not in name_sizes]
    name_sizes += [s for s in range(26, 15, -1) if name_lines(s) == 3 and s not in name_sizes]
    titles: list[str | None] = [title] if title else [None]
    if title and _short_title(title) != title:
        titles.append(_short_title(title))

    def first_fit(title_choices: list[str | None]) -> list[Para] | None:
        for title_text in title_choices:
            for name_size in name_sizes:
                for title_size in (14, 13, 12, 11):
                    paras = build(name_size, title_size, title_text)
                    if title_text and line_count(paras[1], nw) > 2:
                        continue
                    if box_height(paras, nw) <= available:
                        if title_text != title:
                            notes.append(f"Titre raccourci pour tenir dans l'en-tête : « {title_text} »")
                        return paras
        return None

    chosen = first_fit(titles)
    if chosen is None and title:
        # Dernier recours : titre tronqué au mot, avec « … » visible (rien n'est réécrit).
        words = _short_title(title).split()
        while len(words) > 1 and chosen is None:
            words.pop()
            chosen = first_fit([" ".join(words) + " …"])
    if chosen is None:
        chosen = build(name_sizes[-1] if name_sizes else 22, 11, None)
        notes.append("Nom très long : vérifier l'en-tête")
    set_geom(name_box, x=nx, y=top, w=nw, h=available)
    set_text(name_box, chosen, anchor="b")

    if pill is None:
        return
    if not label:
        remove(pill)
        return
    _, _, pw, _ = get_geom(pill)
    px = nx + TEXT_INSET  # bord gauche de la pastille aligné sur le texte du nom
    width = max(pw, text_width(label, "regular", 12) / PT_PER_CM + 2 * TEXT_INSET + 0.8)
    set_geom(pill, x=px, w=width)
    set_text(pill, [Para([Run(label, "regular", 12, WHITE)], align="ctr")], anchor="ctr")


def _clean_linkedin(value: str) -> str:
    """URL LinkedIn sans paramètres de suivi (les exports LinkedIn ajoutent « ?jobid=…&lipi=… »)."""
    return re.split(r"[?#]", value.strip(), maxsplit=1)[0]


def _linkedin_display(value: str) -> str:
    value = re.sub(r"^https?://", "", _clean_linkedin(value), flags=re.I)
    value = re.sub(r"^www\.", "", value, flags=re.I)
    return value.rstrip("/")


def _contact_url(name: str, value: str) -> str | None:
    if name == "email" and "@" in value:
        return "mailto:" + value.strip()
    if name == "linkedin":
        url = _clean_linkedin(value)
        return url if re.match(r"https?://", url, re.I) else "https://" + url
    return None


def _add_hyperlink(t: Template, box: Element, url: str) -> None:
    """Rend la zone cliquable (lien conservé dans le PDF exporté). Le lien est posé sur la forme et non
    sur le texte : PowerPoint soulignerait sinon le texte, contrairement au style du modèle."""
    rel_id = t.slide.part.relate_to(url, RT.HYPERLINK, is_external=True)
    c_nv_pr = child(box, "p:nvSpPr/p:cNvPr")
    for old in c_nv_pr.findall("a:hlinkClick", NS):
        c_nv_pr.remove(old)
    link = etree.Element(q("a:hlinkClick"))
    link.set(q("r:id"), rel_id)
    c_nv_pr.insert(0, link)


def _place_contacts(t: Template, values: dict[str, str | None]) -> None:
    slots = sorted(get_geom(t.placeholders[n])[1] for n in CONTACT_FIELDS if n in t.placeholders)
    slot = 0
    for name in CONTACT_FIELDS:
        box = t.placeholders.get(name)
        if box is None:
            continue
        icon = t.icons.get(name)
        value = values.get(name)
        if not value:
            remove(box)
            remove(icon)
            continue
        if name == "linkedin":
            value = _linkedin_display(value)
        x, y, _, _ = get_geom(box)
        new_y = slots[slot]
        slot += 1
        width = CONTACT_RIGHT_EDGE - x
        size = 10.0
        while size > 7 and not fits_one_line(value, "regular", size, width):
            size -= 0.5
        set_geom(box, y=new_y, w=width)
        set_text(box, [Para([Run(value, "regular", size, BLACK)])], anchor="ctr")
        url = _contact_url(name, values.get(name) or value)
        if url:
            _add_hyperlink(t, box, url)
        if icon is not None:
            _, iy, _, _ = get_geom(icon)
            set_geom(icon, y=iy + (new_y - y))


# --- Colonne de droite --------------------------------------------------------------


RIGHT_MINIMUMS = {"si": 5, "certifications": 2, "expertise": 3, "formations": 1, "langues": 1}
RIGHT_NAMES = {"si": "SI & outils", "certifications": "certifications", "expertise": "expertise", "formations": "formations", "langues": "langues"}


Item = str | Langue | Certification | Formation
ItemList = list[str] | list[Langue] | list[Certification] | list[Formation]


def _item_label(item: Item) -> str:
    if isinstance(item, Certification):
        return item.intitule
    if isinstance(item, Formation):
        return item.diplome
    if isinstance(item, Langue):
        return item.langue
    return item


@dataclass
class _Placement:
    element: Element | None
    y: float
    height: float = 0.0
    paras: list[Para] | None = None  # None : titre de section (seule sa position change)


Blocks = list[tuple[Element, list[Para]]]  # (zone du modèle, paragraphes) des blocs de Compétences
Sections = list[tuple[str, Element, list[Para]]]  # (titre, zone, paragraphes) de Certifications et Formation


def _right_column(t: Template, cv: CV, notes: list[str]) -> None:
    """Empile Compétences / Certifications / Formation sur la page 1, sans jamais créer de page :
    réduction jusqu'à 75 %, puis retrait des éléments les moins prioritaires (signalés au rapport)."""
    expertise, si = list(cv.expertise), list(cv.outils_si)
    langues, certifications, formations = list(cv.langues), list(cv.certifications), list(cv.formations)
    items: dict[str, ItemList] = {"expertise": expertise, "langues": langues, "si": si, "certifications": certifications, "formations": formations}
    for key, label in KEYWORD_BLOCKS.items():
        if 0 < len(items[key]) < MIN_KEYWORD_ITEMS:
            listed = ", ".join(_item_label(i) for i in items[key])
            notes.append(f"Bloc « {label} » non affiché : seulement {len(items[key])} élément(s) ({listed}), trop peu pour un bloc")
            items[key].clear()
    dropped: dict[str, list[str]] = {k: [] for k in items}
    details = True
    comp_header = t.headers.get("Compétences")
    top = get_geom(comp_header)[1] if comp_header is not None else 6.93
    boxes = {key: t.placeholders.get(key) for key, _ in RIGHT_BLOCKS}
    section_boxes = {"Certifications": t.placeholders.get("certifications"), "Formation": t.placeholders.get("formation")}
    builders: dict[str, Callable[[], list[Para]]] = {
        "expertise": lambda: contenu.bullet_paras(expertise),
        "langues": lambda: contenu.langues_paras(langues),
        "si": lambda: contenu.inline_para(si),
    }

    def content() -> tuple[Blocks, Sections]:
        comp: Blocks = []
        for key, label in RIGHT_BLOCKS:
            box, paras = boxes.get(key), builders[key]()
            if box is not None and paras:
                comp.append((box, [contenu.label_para(label)] + paras))
        sections: Sections = []
        for title, paras in (
            ("Certifications", contenu.certification_paras(certifications)),
            ("Formation", contenu.formation_paras(formations, details)),
        ):
            box = section_boxes.get(title)
            if paras and box is not None:
                sections.append((title, box, paras))
        return comp, sections

    def stack(factor: float, comp: Blocks, sections: Sections) -> tuple[float, list[_Placement]]:
        placements: list[_Placement] = []
        y = top
        if comp:
            placements.append(_Placement(comp_header, y))
            y += SECTION_HEADER_HEIGHT - 0.1
            for box, paras in comp:
                scaled = scale(paras, factor)
                h = box_height(scaled, RIGHT_WIDTH)
                placements.append(_Placement(box, y, h, scaled))
                y += h - 0.12
            y += 0.3
        for title, box, paras in sections:
            placements.append(_Placement(t.headers.get(title), y))
            y += SECTION_HEADER_HEIGHT
            scaled = scale(paras, factor)
            h = box_height(scaled, RIGHT_WIDTH)
            placements.append(_Placement(box, y, h, scaled))
            y += h + 0.3
        return y - 0.3, placements

    while True:
        comp, sections = content()
        for factor in (1.0, 0.95, 0.9, 0.85, 0.8, 0.75):
            bottom, placements = stack(factor, comp, sections)
            if bottom <= CONTENT_BOTTOM:
                break
        if bottom <= CONTENT_BOTTOM:
            break
        if details and any(f.details for f in formations):
            details = False
            notes.append("Colonne de droite : spécialités/options des formations masquées faute de place")
            continue
        excess = [(len(items[k]) - minimum, k) for k, minimum in RIGHT_MINIMUMS.items() if len(items[k]) > minimum]
        if not excess:
            notes.append("Colonne de droite très chargée : vérifier le bas de la page 1")
            break
        key = max(excess, key=lambda e: e[0])[1]
        dropped[key].insert(0, _item_label(items[key].pop()))

    lost = [f"{RIGHT_NAMES[k]} : {', '.join(v)}" for k, v in dropped.items() if v]
    if lost:
        notes.append("Faute de place sur la page, non affichés (conservés dans le JSON) — " + " ; ".join(lost))
    if factor < 1.0:
        notes.append(f"Colonne de droite réduite à {int(factor * 100)} % pour tenir sur la page")
    used = {id(box) for box, _ in comp} | {id(box) for _, box, _ in sections}
    for box in boxes.values():
        if box is not None and id(box) not in used:
            remove(box)
    for title, box in section_boxes.items():
        if box is not None and id(box) not in used:
            remove(box)
            remove(t.headers.get(title))
    if not comp:
        remove(comp_header)
    for placement in placements:
        if placement.element is None:
            continue
        if placement.paras is None:  # titre de section
            set_geom(placement.element, y=placement.y)
        else:
            set_geom(placement.element, x=RIGHT_X, y=placement.y, w=RIGHT_WIDTH, h=placement.height)
            set_text(placement.element, placement.paras)


# --- Flux des expériences (colonne gauche puis pages de suite) -------------------------


@dataclass
class Header:
    title: str
    source: str  # titre du groupe du modèle à copier


@dataclass
class Segment:
    kind: str  # "text" | "header"
    y: float
    paras: list[Para] = field(default_factory=list[Para])
    title: str = ""
    source: str = ""


@dataclass
class Page:
    x: float
    top: float
    width: float
    segments: list[Segment] = field(default_factory=list[Segment])


MAX_PARA_HEIGHT = 12.0  # cm : au-delà, un paragraphe est coupé en morceaux paginables


def split_long_para(para: Para, width: float) -> list[Para]:
    """Coupe un paragraphe démesuré (ex. une puce de 3 000 caractères) en morceaux d'au plus
    MAX_PARA_HEIGHT, au mot près, sans rien retirer."""
    if para_height(para, width) <= MAX_PARA_HEIGHT:
        return [para]
    words = [(piece, run) for run in para.runs for piece in run.text.split(" ") if piece]

    def make(chunk: list[tuple[str, Run]], first: bool) -> Para:
        runs: list[Run] = []
        for piece, run in chunk:
            if runs and runs[-1].style == run.style and runs[-1].size == run.size and runs[-1].color == run.color:
                runs[-1] = Run(runs[-1].text + " " + piece, run.style, run.size, run.color)
            else:
                runs.append(Run((" " if runs else "") + piece, run.style, run.size, run.color))
        return dataclasses.replace(
            para,
            runs=runs,
            bullet=para.bullet if first else None,
            space_before=para.space_before if first else 0,
            keep_with_next=False,
        )

    chunks: list[Para] = []
    start, first = 0, True
    while start < len(words):
        lo, hi = start + 1, len(words)
        while lo < hi:  # plus grand préfixe qui tient dans la hauteur maximale
            mid = (lo + hi + 1) // 2
            if para_height(make(words[start:mid], first), width) <= MAX_PARA_HEIGHT:
                lo = mid
            else:
                hi = mid - 1
        chunks.append(make(words[start:lo], first))
        start, first = lo, False
    return chunks


def _units(paras: list[Para]) -> list[list[Para]]:
    """Regroupe les paragraphes « garder avec le suivant » avec leur suivant."""
    units: list[list[Para]] = []
    current: list[Para] = []
    for para in (piece for p in paras for piece in split_long_para(p, LEFT_WIDTH)):
        current.append(para)
        if not para.keep_with_next:
            units.append(current)
            current = []
    if current:
        units.append(current)
    return units


@dataclass
class _FlowState:
    text: Segment | None = None  # zone de texte en cours de remplissage


def _paginate(flow: list[Block | Header], has_experience_header: bool) -> list[Page]:
    pages = [Page(LEFT_X, LEFT_TOP_PAGE1, LEFT_WIDTH)]
    # Section en cours : (titre affiché, groupe du modèle à copier pour l'en-tête « (suite) »).
    section: tuple[str, str] | None = ("Expériences", "Expériences") if has_experience_header else None
    state = _FlowState()

    def page() -> Page:
        return pages[-1]

    def text_seg(y: float) -> Segment:
        seg = Segment("text", y)
        page().segments.append(seg)
        state.text = seg
        return seg

    def cursor() -> float:
        segs = page().segments
        if not segs:
            return page().top
        last = segs[-1]
        if last.kind == "header":
            return last.y + SECTION_HEADER_HEIGHT
        return last.y + (box_height(last.paras, page().width) if last.paras else 0)

    def new_page(suite: tuple[str, str] | None, continuation: Para | None) -> Segment:
        pages.append(Page(LEFT_X, CONTINUATION_HEADER_Y, FULL_WIDTH))
        if suite:
            page().segments.append(Segment("header", page().top, title=f"{suite[0]} (suite)", source=suite[1]))
        seg = text_seg(cursor())
        if continuation is not None:
            seg.paras.append(continuation)
        return seg

    if not has_experience_header:
        page().top = 6.93
    for item in flow:
        if isinstance(item, Header):
            y = cursor()
            if page().segments:
                y += HEADER_GAP
            if y + SECTION_HEADER_HEIGHT + 3.0 > CONTENT_BOTTOM:  # pas de titre orphelin en bas de page
                new_page(None, None)
                page().segments.clear()
                y = page().top
            page().segments.append(Segment("header", y, title=item.title, source=item.source))
            section = (item.title, item.source)
            text_seg(cursor())
            continue
        block: Block = item
        seg = state.text if state.text is not None and state.text is page().segments[-1] else text_seg(cursor())
        for index, unit in enumerate(_units(block.paras)):
            trial = seg.paras + unit
            if seg.y + box_height(trial, page().width) <= CONTENT_BOTTOM or not seg.paras:
                seg.paras.extend(unit)
                continue
            continuation = block.continuation if index > 0 else None
            seg = new_page(section, continuation)
            seg.paras.extend(unit)
    for p in pages:
        p.segments = [s for s in p.segments if s.kind == "header" or s.paras]
    return pages


def _scale_block(block: Block, factor: float) -> Block:
    if factor == 1.0:
        return block
    continuation = scale([block.continuation], factor)[0] if block.continuation else None
    return Block(scale(block.paras, factor), continuation)


def _split_experiences(cv: CV) -> tuple[list[Experience], list[Experience]]:
    """Expériences professionnelles (antichronologiques) et bénévolat (relégué en fin de flux)."""
    experiences = contenu.sort_experiences(cv.experiences)
    return [e for e in experiences if e.type != "benevolat"], [e for e in experiences if e.type == "benevolat"]


def _build_flow(
    cv: CV,
    factor: float = 1.0,
    condensed: frozenset[int] = frozenset(),
    projects: str = "full",
    volunteering: bool = True,
) -> tuple[list[Block | Header], bool]:
    """Flux de la colonne gauche : expériences, projets, puis bénévolat. Rien d'autre ne crée de page."""
    pro, benevolat = _split_experiences(cv)
    flow: list[Block | Header] = [_scale_block(contenu.experience_block(exp, condensed=i in condensed), factor) for i, exp in enumerate(pro)]
    if cv.projets and projects != "none":
        flow.append(Header("Projets", "Expériences"))
        flow += [_scale_block(contenu.project_block(p, titles_only=projects == "titles"), factor) for p in cv.projets]
    if benevolat and volunteering:
        flow.append(Header("Bénévolat", "Expériences"))
        flow += [_scale_block(contenu.experience_block(exp), factor) for exp in benevolat]
    return flow, bool(pro or benevolat)


NEARLY_EMPTY = 0.3  # dernière page remplie à moins de 30 % : on essaie de s'en passer


def _page_fill(page: Page) -> float:
    """Part de la hauteur utile occupée sur une page."""
    if not page.segments:
        return 0.0
    last = page.segments[-1]
    bottom = last.y + (SECTION_HEADER_HEIGHT if last.kind == "header" else box_height(last.paras, page.width))
    return (bottom - page.top) / (CONTENT_BOTTOM - page.top)


def _fit_left(cv: CV, limit: int | None, notes: list[str], fallback: bool = True) -> tuple[list[Page], bool]:
    """Pagination de la colonne gauche.

    - limite fixée (CV d'une page, ou --pages-max) : réduction jusqu'à 80 %, puis bénévolat retiré,
      projets condensés ou retirés, puis puces masquées en partant des expériences les plus
      anciennes (les 2 plus récentes restent complètes) ;
    - sans limite (CV source de plusieurs pages) : les pages de suite ne servent qu'aux expériences
      professionnelles ; bénévolat et projets qui ajouteraient une page sont retirés ou condensés ; une
      dernière page presque vide est évitée en condensant comme pour un CV d'une page (cas des exports
      LinkedIn, qui font plusieurs pages pour peu de contenu).
    Rien n'est réécrit : seuls des détails sont masqués, et le rapport le dit.
    """
    pro, benevolat = _split_experiences(cv)
    has_exp = bool(pro or benevolat)
    extras = bool(cv.projets or benevolat)

    def run(factor: float, condensed: frozenset[int] = frozenset(), projects: str = "full", volunteering: bool = True) -> list[Page]:
        return _paginate(_build_flow(cv, factor, condensed, projects, volunteering)[0], has_exp)

    hidden: list[str] = []
    if limit is None:
        target = len(run(1.0, projects="none", volunteering=False)) if extras else None
        pages, factor, mode, volunteering = run(1.0), 1.0, "full", True
        for candidate in (0.95, 0.9, 0.85):  # évite une dernière page presque vide
            if len(pages) <= 1:
                break
            trial = run(candidate)
            if len(trial) < len(pages):
                pages, factor = trial, candidate
                break
        if target is not None and len(pages) > target and benevolat:
            volunteering = False
            pages = run(factor, volunteering=False)
            hidden.append("bénévolat")
        if target is not None and len(pages) > target and cv.projets:
            for mode in ("titles", "none"):
                pages = run(factor, projects=mode, volunteering=volunteering)
                if len(pages) <= target:
                    break
            hidden.append("détail des projets" if mode == "titles" else "projets")
        if len(pages) > 1 and _page_fill(pages[-1]) < NEARLY_EMPTY:
            trial_notes: list[str] = []
            trial = _fit_left(cv, len(pages) - 1, trial_notes, fallback=False)
            if len(trial[0]) < len(pages):
                notes.extend(trial_notes)
                return trial
        if factor < 1.0:
            notes.append(f"Expériences réduites à {int(factor * 100)} % pour éviter une page presque vide")
    else:
        for factor in (1.0, 0.95, 0.9, 0.85, 0.8):
            pages = run(factor)
            if len(pages) <= limit:
                break
        mode, volunteering = "full", True
        condensed: set[int] = set()
        if len(pages) > limit and benevolat:
            volunteering = False
            pages = run(factor, volunteering=False)
            hidden.append("bénévolat")
        if len(pages) > limit and cv.projets:
            for mode in ("titles", "none"):
                pages = run(factor, projects=mode, volunteering=volunteering)
                if len(pages) <= limit:
                    break
            hidden.append("détail des projets" if mode == "titles" else "projets")
        if len(pages) > limit:
            for index in range(len(pro) - 1, 1, -1):
                if not pro[index].realisations:
                    continue
                condensed.add(index)
                pages = run(factor, frozenset(condensed), mode, volunteering)
                if len(pages) <= limit:
                    break
        if len(pages) > limit:
            if not fallback:  # essai pour éviter une page presque vide : sans succès, rien n'est masqué
                return pages, has_exp
            # Même condensé, le CV ne tient pas : inutile de masquer quoi que ce soit, les expériences
            # continuent sur des pages de suite (cas des profils très expérimentés).
            notes.append(f"Trop d'expériences pour {limit} page(s) : pages de suite réservées aux expériences")
            return _fit_left(cv, None, notes)
        if condensed:
            names = [" — ".join(x for x in (pro[i].poste, pro[i].entreprise) if x) for i in sorted(condensed)]
            hidden.append("réalisations de : " + " ; ".join(names))
        if factor < 1.0:
            notes.append(f"Expériences réduites à {int(factor * 100)} % pour tenir sur {limit} page(s)")
    if hidden:
        notes.append("Faute de place, masqué (conservé dans le JSON) : " + " | ".join(hidden))
    return pages, has_exp


def _render_pages(t: Template, pages: list[Page], has_exp: bool) -> None:
    exp_box = t.placeholders["experiences"]
    exp_header = t.headers.get("Expériences")
    prototype_box = copy.deepcopy(exp_box)
    header_sources = dict(t.headers)
    layout = t.slide.slide_layout

    for index, page in enumerate(pages):
        if index == 0:
            tree = t.tree
        else:
            tree = t.prs.slides.add_slide(layout).shapes.element
        used_template_box = False
        for seg in page.segments:
            if seg.kind == "header":
                if index == 0 and has_exp and seg.title == "Expériences":
                    continue
                source = header_sources.get(seg.source)
                if source is None:  # jamais « a or b » sur un élément lxml : un élément sans enfant est « faux »
                    source = exp_header
                if source is None:
                    continue
                group = t.clone(source, tree)
                _retitle_header(group, seg.title)
                set_geom(group, x=LEFT_X + 0.4 if page.x == LEFT_X else page.x, y=seg.y)
                continue
            if index == 0 and not used_template_box:
                box = exp_box
                used_template_box = True
            else:
                box = t.clone(prototype_box, tree)
            set_geom(box, x=page.x, y=seg.y, w=page.width, h=box_height(seg.paras, page.width))
            set_text(box, seg.paras)
    if not used_template_box_on_first(pages):
        remove(exp_box)
    if not has_exp:
        remove(exp_header)


def used_template_box_on_first(pages: list[Page]) -> bool:
    return any(seg.kind == "text" for seg in pages[0].segments) if pages else False


def _remove_unused_placeholders(t: Template, notes: list[str]) -> None:
    """Zone encore marquée {{…}} (champ que l'outil ne remplit pas, ex. {{methode}} d'un ancien modèle) :
    retirée plutôt que laissée telle quelle dans le CV."""
    for el in [el for el in t.tree if el.tag in SHAPE_TAGS]:
        names = re.findall(r"\{\{\s*(\w+)\s*\}+", text_of(el))
        if names:
            remove(el)
            notes.append(f"Zone du modèle non utilisée retirée : {', '.join('{{' + n + '}}' for n in names)}")


# --- Point d'entrée -----------------------------------------------------------------


@dataclass
class RenderResult:
    pages: int
    notes: list[str]


def render_cv(
    cv: CV,
    photo: Path | None,
    template: Path,
    output: Path,
    *,
    experience_label: str | None,
    contact: dict[str, str | None],
    name: str | None = None,
    document_title: str | None = None,
    pages_max: int | None = None,
    source_pages: int = 1,
) -> RenderResult:
    prs = Presentation(str(template))
    prepare_presentation(prs)
    t = Template(prs)
    notes: list[str] = []

    has_photo = _place_photo(t, photo)
    title = contenu.tame_caps(contenu.clean_item(cv.titre)) if cv.titre else None
    # Un même intitulé rangé à la fois en certification et en formation n'est affiché qu'en formation.
    diplomas = {contenu.strip_accents(f.diplome.lower()).strip() for f in cv.formations}
    certifications = [c for c in cv.certifications if contenu.strip_accents(c.intitule.lower()).strip() not in diplomas]
    if len(certifications) != len(cv.certifications):
        cv = cv.model_copy(update={"certifications": certifications})
    _place_identity(t, name or display_name(cv), title, experience_label, has_photo, notes)
    _place_contacts(t, contact)
    _right_column(t, cv, notes)
    # Une page par défaut ; des pages de suite (expériences uniquement) seulement si le CV
    # d'origine en comptait déjà plusieurs, ou si --pages-max le demande explicitement.
    limit = pages_max if pages_max else (1 if source_pages <= 1 else None)
    pages, has_exp = _fit_left(cv, limit, notes)
    _render_pages(t, pages, has_exp)
    _remove_unused_placeholders(t, notes)

    props = prs.core_properties
    props.title = document_title or f"CV Logiclever - {name or display_name(cv)}"
    props.author = "Logiclever"
    props.last_modified_by = "Logiclever"
    props.subject = ""
    props.keywords = ""
    props.comments = ""
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output))
    return RenderResult(len(pages), notes)

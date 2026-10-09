"""Calibrage du rendu PDF : PowerPoint (Windows) et LibreOffice (Linux, Cloud Run).

À relancer quand PowerPoint, LibreOffice ou les polices changent de version (voir README, « Rendu sous Linux »).

  python outils/calibrage_libreoffice.py diapositive calibrage.pptx
      diapositives de calibrage (code de production : mêmes zones, paragraphes et styles que les CV)
      + calibrage.json (géométrie et contenu de chaque zone)
  python outils/calibrage_libreoffice.py mesurer calibrage.json calibrage.pdf
      lignes de base mesurées dans le PDF rendu par PowerPoint ou LibreOffice -> calibrage.mesures.json
  python outils/calibrage_libreoffice.py facteurs calibrage.json calibrage.mesures.json
      position de la ligne de base PowerPoint selon l'interligne : valeurs de config.POWERPOINT_BASELINE
  python outils/calibrage_libreoffice.py comparer DOSSIER_PDF_POWERPOINT DOSSIER_PDF_LIBREOFFICE
      CV par CV : lignes identiques, écarts des lignes de base et des puces (mm)

Pour obtenir les PDF de référence sous Windows : python -m cv_formatter --depuis-json --sortie ref (PowerPoint) ;
sous Linux : même commande dans l'image Docker (LibreOffice calé). Le contrôle automatique sans PowerPoint est
tests/test_rendu_linux.py.
"""

from __future__ import annotations

import copy
import difflib
import json
import re
import statistics as st
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import pymupdf  # noqa: E402
from pptx import Presentation  # noqa: E402

from cv_formatter.config import TEMPLATE_PATH, TEXT_INSET  # noqa: E402
from cv_formatter.layout import PT_PER_CM, Para, Run, box_height  # noqa: E402
from cv_formatter.render_pptx import SHAPE_TAGS, Template, prepare_presentation, remove, set_geom, set_text  # noqa: E402

SIZES = [6.5, 7, 7.5, 8, 8.5, 9, 9.5, 10, 10.5, 11, 11.5, 12]
# Interligne -> tailles où il est employé : 0,9 ne sert qu'au nom du consultant (16 à 37 pt).
SPACINGS = {0.9: [16, 20, 24, 28, 32, 37], 1.0: SIZES + [14, 18], 1.05: SIZES, 1.1: SIZES, 1.15: SIZES}
WIDTH = 6.5  # cm : aucune ligne de calibrage ne doit passer à la ligne
MM = 10 / PT_PER_CM  # points -> mm


def _zones() -> list[dict]:
    zones: list[dict] = []
    for ls, sizes in SPACINGS.items():  # interligne et première ligne : 4 paragraphes d'une ligne
        for size in sizes:
            zones.append({"type": "interligne", "paras": [{"texte": f"l{i}", "style": "regular", "taille": size, "ls": ls, "avant": 0} for i in range(4)]})
    for avant in (1, 2, 3, 4, 5, 10, 11):  # espace avant (points entiers)
        for size in (8, 9, 10, 12):
            zones.append({"type": "espace", "paras": [{"texte": f"espace {i}", "style": "regular", "taille": size, "ls": 1.1, "avant": avant} for i in range(4)]})
    for style in ("regular", "medium", "semibold", "bold"):  # graisses
        for size in (8, 10, 12):
            zones.append({"type": "graisse", "paras": [{"texte": f"{style} {i}", "style": style, "taille": size, "ls": 1.0, "avant": 0} for i in range(3)]})
    for size in (7.5, 8, 9, 10):  # puces (Arial)
        zones.append({"type": "puce", "paras": [{"texte": f"puce {i}", "style": "regular", "taille": size, "ls": 1.1, "avant": 2, "puce": True} for i in range(4)]})
    return zones


def diapositive(sortie: Path) -> None:
    prs = Presentation(str(TEMPLATE_PATH))
    prepare_presentation(prs)
    t = Template(prs)
    prototype = copy.deepcopy(t.placeholders["experiences"])
    layout = t.slide.slide_layout
    tree = t.tree
    for el in [el for el in tree if el.tag in SHAPE_TAGS]:
        remove(el)
    zones, page, x, y, colonne = _zones(), 0, 0.5, 0.5, WIDTH
    for numero, zone in enumerate(zones):
        zone["id"] = numero
        paras = [Para([Run(f"Z{numero:03d} Hxgé {p['texte']}", p["style"], p["taille"])], bullet="•" if p.get("puce") else None,
                      indent=0.4 if p.get("puce") else 0.0, space_before=p["avant"], line_spacing=p["ls"]) for p in zone["paras"]]
        largeur = WIDTH if zone["paras"][0]["taille"] <= 14 else 3 * WIDTH  # grands corps (nom) : zone plus large
        hauteur = box_height(paras, largeur) + 0.1
        if y + hauteur > 28.5:
            x, y = x + colonne + 0.3, 0.5
        if x + largeur > 20.8:
            tree = prs.slides.add_slide(layout).shapes.element
            for el in [el for el in tree if el.tag in SHAPE_TAGS]:
                remove(el)
            page, x, y = page + 1, 0.5, 0.5
        colonne = largeur if y == 0.5 else max(colonne, largeur)
        zone_xml = t.clone(prototype, tree)
        set_geom(zone_xml, x=x, y=y, w=largeur, h=hauteur)
        set_text(zone_xml, paras)
        zone.update({"page": page, "x": x, "y": y})
        y += hauteur + 0.25
    prs.save(str(sortie))
    sortie.with_suffix(".json").write_text(json.dumps(zones, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"{len(zones)} zones sur {page + 1} diapositives : {sortie} (+ {sortie.with_suffix('.json').name})")
    print("Rendre ce fichier en PDF avec PowerPoint (Enregistrer sous > PDF) et/ou LibreOffice, puis « mesurer ».")


def mesurer(spec: Path, pdf: Path) -> None:
    zones = json.loads(spec.read_text(encoding="utf-8"))
    lignes: dict[int, list[dict]] = {}
    with pymupdf.open(pdf) as doc:
        for numero_page, page in enumerate(doc):
            for block in page.get_text("dict")["blocks"]:
                for line in block.get("lines", []):
                    spans = [s for s in line["spans"] if s["text"].strip() and s["size"] > 2]
                    texte = "".join(s["text"] for s in spans)
                    trouve = re.search(r"Z(\d{3}) ", texte)
                    if trouve:
                        texte_spans = [s for s in spans if s["text"].strip() != "•"]
                        puce = next((s for s in spans if s["text"].strip() == "•"), None)
                        lignes.setdefault(int(trouve[1]), []).append({
                            "page": numero_page, "base": texte_spans[0]["origin"][1] / PT_PER_CM,
                            "puce": puce["origin"][1] / PT_PER_CM if puce else None})
    mesures = [{"id": z["id"], "lignes": sorted(lignes.get(z["id"], []), key=lambda l: l["base"])} for z in zones]
    cible = pdf.with_suffix(".mesures.json")
    cible.write_text(json.dumps(mesures, indent=1), encoding="utf-8")
    print(f"{sum(len(m['lignes']) for m in mesures)} lignes mesurées -> {cible}")


def facteurs(spec: Path, mesures_path: Path) -> None:
    zones = {z["id"]: z for z in json.loads(spec.read_text(encoding="utf-8"))}
    par_interligne: dict[float, list[float]] = {}
    ecarts_pas: list[float] = []
    for mesure in json.loads(mesures_path.read_text(encoding="utf-8")):
        zone, lignes = zones[mesure["id"]], mesure["lignes"]
        if len(lignes) != len(zone["paras"]):  # ligne coupée en deux : zone ignorée
            continue
        p = zone["paras"][0]
        premiere = (lignes[0]["base"] - zone["y"] - TEXT_INSET) * PT_PER_CM / p["taille"]
        par_interligne.setdefault(p["ls"], []).append(premiere)
        attendu = 1.2 * p["taille"] * p["ls"] + p["avant"]
        ecarts_pas += [((b["base"] - a["base"]) * PT_PER_CM - attendu) * MM for a, b in zip(lignes, lignes[1:])]
    print("Ligne de base PowerPoint (em depuis le haut de la ligne) — à reporter dans config.POWERPOINT_BASELINE :")
    print("POWERPOINT_BASELINE = {" + ", ".join(f"{ls}: {st.mean(v):.4f}" for ls, v in sorted(par_interligne.items())) + "}")
    for ls, valeurs in sorted(par_interligne.items()):
        print(f"  interligne {ls} : {st.mean(valeurs):.4f} em (écart-type {st.pstdev(valeurs):.4f}, {len(valeurs)} zones)")
    print(f"Pas entre lignes, écart au modèle 1,2 × taille × interligne : moyen {st.mean(ecarts_pas):+.4f} mm, max {max(map(abs, ecarts_pas)):.4f} mm")


def _lignes_pdf(pdf: Path) -> list[dict]:
    out = []
    with pymupdf.open(pdf) as doc:
        for numero_page, page in enumerate(doc):
            for block in page.get_text("rawdict")["blocks"]:
                for line in block.get("lines", []):
                    spans = [s for s in line["spans"] if s["size"] > 2 and "".join(c["c"] for c in s["chars"]).strip()]
                    if not spans:
                        continue
                    puce = spans[0] if "".join(c["c"] for c in spans[0]["chars"]).strip() == "•" else None
                    texte = spans[1:] if puce else spans
                    if texte:
                        out.append({"page": numero_page, "texte": re.sub(r"\s+", "", "".join(c["c"] for s in texte for c in s["chars"])),
                                    "base": texte[0]["origin"][1], "puce": puce["origin"][1] if puce else None})
    return out


def comparer(reference: Path, essai: Path) -> None:
    print(f"{'CV':30} {'lignes':>9} {'ident.':>6} {'dy moy':>7} {'dy max':>7} {'puce':>6}  (mm ; césure différente : « dy max » d'une ligne ou plus)")
    tous: list[float] = []
    for ref in sorted(reference.glob("*.pdf")):
        test = essai / ref.name
        if not test.exists():
            print(f"{ref.stem[:30]:30} absent de {essai}")
            continue
        a, b = _lignes_pdf(ref), _lignes_pdf(test)
        sm = difflib.SequenceMatcher(a=[(l["page"], l["texte"]) for l in a], b=[(l["page"], l["texte"]) for l in b], autojunk=False)
        dy, puces = [], []
        for op, i1, i2, j1, j2 in sm.get_opcodes():
            if op == "equal":
                for i, j in zip(range(i1, i2), range(j1, j2)):
                    dy.append((b[j]["base"] - a[i]["base"]) * MM)
                    if a[i]["puce"] is not None and b[j]["puce"] is not None:
                        puces.append(abs(b[j]["puce"] - a[i]["puce"]) * MM)
        tous += dy
        print(f"{ref.stem.replace('CV Logiclever - ', '')[:30]:30} {len(a):>4}/{len(b):<4} {len(dy):>6} {st.mean(dy):+7.3f} {max(map(abs, dy)):7.3f} {max(puces, default=0):6.3f}")
    if tous:
        print(f"Ensemble : {len(tous)} lignes identiques, écart moyen {st.mean(tous):+.3f} mm, médiane des écarts absolus {st.median(map(abs, tous)):.3f} mm")


if __name__ == "__main__":
    commandes = {"diapositive": (diapositive, 1), "mesurer": (mesurer, 2), "facteurs": (facteurs, 2), "comparer": (comparer, 2)}
    if len(sys.argv) < 2 or sys.argv[1] not in commandes or len(sys.argv) != 2 + commandes[sys.argv[1]][1]:
        print(__doc__)
        sys.exit(2)
    fonction, _ = commandes[sys.argv[1]]
    fonction(*(Path(a) for a in sys.argv[2:]))

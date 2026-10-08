"""Tests de non-régression (données fictives uniquement : aucun vrai CV dans le dépôt).

Lancer : python -m pytest -q
"""

from __future__ import annotations

import io
import zipfile
from collections.abc import Iterable, Iterator
from pathlib import Path
from types import SimpleNamespace as NS
from typing import Any

import pymupdf
import pytest
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.presentation import Presentation as PresentationFile
from pptx.shapes.autoshape import Shape
from pptx.shapes.base import BaseShape
from pptx.shapes.group import GroupShape
from pptx.util import Emu

from cv_formatter import contenu
from cv_formatter.__main__ import Job, collect_files, default_input, json_edited, needs_extraction, photo_decision, safe_name
from cv_formatter.anonymisation import anonymize, initials
from cv_formatter.config import (
    FONTS_DIR,
    HEADER_TEXT_OFFSET,
    INPUT_DIR,
    LEFT_X,
    RIGHT_X,
    TEMPLATE_PATH,
    TEXT_INSET,
    WIDTH_SAFETY_MARGIN,
    cm,
)
from cv_formatter.controle import apply_verdicts, build_claims, fingerprint_data, photo_verdict
from cv_formatter.extraction import EXTRACTION_VERSION
from cv_formatter.layout import PT_PER_CM, Para, Run, cap_center_offset, line_count, para_height, scale, text_width
from cv_formatter.pdf_source import SourceDocument, fix_text, load_source
from cv_formatter.rapport import CVReport
from cv_formatter.render_pptx import NS as XML_NS, RenderResult, render_cv, split_long_para, text_of
from cv_formatter.schema import CV, Verdict
from cv_formatter.verification import verify
from cv_formatter.visage import FaceCheck, check_face

SOURCE = """Jean DUPONT
Chef de projet SI
jean.dupont@example.com | 06 12 34 56 78 | Lyon (69003)
linkedin.com/in/jean-dupont
Expériences
Chef de projet SI — ACME Conseil (mission chez Banque Exemple) — Janvier 2020 – Présent
Pilotage d'un portefeuille de 200 projets IT avec Power BI et SAP S4/HANA
Animation des comités de pilotage et suivi du budget de 2,3 M€
Développeur — ACME Conseil — 2016 - 2019 (alternance)
Développement de l'outil de suivi de140 magasins
Compétences : pilotage de portefeuille, gestion de projet, conduite du changement
Outils : Power BI, SAP S4/HANA, Jira
Formation
Master Management de projet — Université Exemple — 2014 - 2016
Certification Scrum PSPO I
Langues : Anglais courant (TOEIC 890), depuis trois ans en gestion de projet
INWI (3ième opérateur telecom)
"""


def make_cv(**overrides: Any) -> CV:
    data = {
        "langue_source": "fr",
        "prenom": "Jean",
        "nom": "DUPONT",
        "titre": "Chef de projet SI",
        "contact": {"email": "jean.dupont@example.com", "telephone": "06 12 34 56 78", "localisation": "Lyon (69003)", "linkedin": "linkedin.com/in/jean-dupont"},
        "annees_experience": {"valeur": None, "citation": None},
        "expertise": ["Pilotage de portefeuille", "Gestion de projet", "Conduite du changement"],
        "outils_si": ["Power BI", "SAP S4/HANA", "Jira"],
        "competences_detaillees": [],
        "langues": [{"langue": "Anglais", "niveau": "Courant (TOEIC 890)"}],
        "certifications": [{"intitule": "Scrum PSPO I", "organisme": None, "date": None}],
        "formations": [{"diplome": "Master Management de projet", "etablissement": "Université Exemple", "lieu": None, "periode": "2014 - 2016", "details": None}],
        "experiences": [
            {"poste": "Chef de projet SI", "entreprise": "ACME Conseil", "client": "Banque Exemple", "lieu": None,
             "periode_texte": "Janvier 2020 – Présent", "debut": "2020-01", "fin": "present", "type": "mission", "contexte": None,
             "realisations": ["Pilotage d'un portefeuille de 200 projets IT avec Power BI et SAP S4/HANA",
                              "Animation des comités de pilotage et suivi du budget de 2,3 M€"], "environnement": ["Power BI", "SAP S4/HANA"]},
            {"poste": "Développeur", "entreprise": "ACME Conseil", "client": None, "lieu": None, "periode_texte": "2016 - 2019",
             "debut": "2016", "fin": "2019", "type": "alternance", "contexte": None,
             "realisations": ["Développement de l'outil de suivi de 140 magasins"], "environnement": []},
        ],
        "projets": [],
        "photo_candidate": None,
        "sections_non_reprises": [],
    }
    data.update(overrides)
    return CV.model_validate(data)


def all_confirmed(cv: CV) -> list[Verdict]:
    return [Verdict(id=c.id, statut="confirme", explication="") for c in build_claims(cv)]


# --- Contrôle déterministe -----------------------------------------------------------------


def test_faithful_extraction_is_untouched():
    checked, findings = verify(make_cv(), SOURCE, [], apply_removals=True)
    assert [f for f in findings if f.action == "retiré"] == []
    assert len(checked.experiences[0].realisations) == 2


@pytest.mark.parametrize(
    "bullet",
    [
        "Pilotage d'un portefeuille de 350 projets IT",  # nombre inventé
        "Certification PMP obtenue",  # sigle inventé
        "Migration vers Salesforce",  # outil inventé (nom propre)
    ],
)
def test_invented_bullet_is_removed(bullet: str):
    cv = make_cv()
    cv.experiences[0].realisations.append(bullet)
    checked, findings = verify(cv, SOURCE, [], apply_removals=True)
    assert bullet not in checked.experiences[0].realisations
    assert any(f.action == "retiré" and f.text == bullet for f in findings)


def test_tolerances_keep_true_information():
    cv = make_cv(outils_si=["PowerBI", "SAP S4/HANA"])  # forme collée
    cv.experiences[1].realisations = ["Développement de l'outil de suivi de 140 magasins"]  # « de140 » dans le PDF
    cv.experiences[0].entreprise = "INWI (3ème opérateur télécom)"  # « 3ième » dans le PDF
    _, findings = verify(cv, SOURCE + "\nACME Conseil", [], apply_removals=True)
    assert [f for f in findings if f.action == "retiré"] == []


def test_contact_must_be_visible():
    cv = make_cv(contact={"email": "autre.personne@example.com", "telephone": "07 99 99 99 99", "localisation": None, "linkedin": None})
    checked, _ = verify(cv, SOURCE, [], apply_removals=True)
    assert checked.contact.email is None and checked.contact.telephone is None


def test_doubtful_reference_only_flags():
    cv = make_cv()
    cv.experiences[0].realisations.append("Pilotage d'un portefeuille de 350 projets IT")
    checked, findings = verify(cv, SOURCE, [], apply_removals=False)
    assert len(checked.experiences[0].realisations) == 3
    assert all(f.action == "à vérifier" for f in findings)


def test_company_name_checked_including_first_word():
    cv = make_cv()
    cv.experiences[1].entreprise = "Capgemini"  # un seul mot, absent du CV : n'était pas contrôlé avant
    checked, findings = verify(cv, SOURCE, [], apply_removals=True)
    assert checked.experiences[1].entreprise is None
    assert any(f.text == "Capgemini" and f.action == "retiré" for f in findings)


# --- Employeur / client ----------------------------------------------------------------------

HEADER = "Mohammed AMMARA\nConsultant confirmé Logiclever\n3 ans d’expérience\nmohammed@logiclever.com\n"
EXPERIENCE = "Clem’ (Opérateur d’infrastructure de recharge)\nProduct Owner : portail web\nAnimer les cérémonies agiles Scrum\n"


def _ammara(**exp: Any) -> CV:
    experience = {"poste": "Product Owner", "entreprise": "Logiclever", "client": "Clem’", "lieu": None, "periode_texte": None,
                  "debut": None, "fin": None, "type": "mission", "contexte": None,
                  "realisations": ["Animer les cérémonies agiles Scrum"], "environnement": []}
    experience.update(exp)
    return make_cv(prenom="Mohammed", nom="AMMARA", titre="Consultant confirmé Logiclever", contact={"email": None, "telephone": None, "localisation": None, "linkedin": None},
                   expertise=[], outils_si=[], langues=[], certifications=[], formations=[], experiences=[experience])


def test_employer_named_only_in_cv_header_is_detached():
    """Cas réel : CV mis en forme par Logiclever pour un consultant qui travaillait chez Clem' AVANT de rejoindre
    Logiclever. « Logiclever » n'apparaît que dans le titre : ce n'est pas l'employeur de l'expérience."""
    body = EXPERIENCE  # texte hors en-tête (cf. SourceDocument.body_text)
    checked, findings = verify(_ammara(), HEADER + EXPERIENCE, [], apply_removals=True, body_text=body)
    exp = checked.experiences[0]
    assert (exp.entreprise, exp.client) == ("Clem’", None)
    assert any(f.text == "Logiclever" and f.action == "retiré" and "en-tête" in f.reason for f in findings)


def test_employer_named_only_in_email_domain_is_detached():
    source = "Mohammed AMMARA\ncontact : mohammed@logiclever.com\n" + EXPERIENCE
    body = EXPERIENCE + "contact :  \n"  # e-mails retirés du texte de référence
    checked, _ = verify(_ammara(), source, [], apply_removals=True, body_text=body)
    assert checked.experiences[0].entreprise == "Clem’"


def test_employer_named_in_experience_or_group_heading_is_kept():
    body = "Logiclever (2022 - Présent)\n" + EXPERIENCE  # titre de rubrique regroupant les missions
    checked, findings = verify(_ammara(), HEADER + body, [], apply_removals=True, body_text=body)
    exp = checked.experiences[0]
    assert (exp.entreprise, exp.client) == ("Logiclever", "Clem’") and not [f for f in findings if f.action == "retiré"]


def test_client_named_only_in_header_is_removed():
    source = "Jean DUPONT\nConsultant chez Banque Exemple\n" + SOURCE.split("\n", 2)[2].replace("(mission chez Banque Exemple) ", "")
    body = source.split("\n", 2)[2]
    checked, _ = verify(make_cv(), source, [], apply_removals=True, body_text=body)
    assert checked.experiences[0].client is None and checked.experiences[0].entreprise == "ACME Conseil"


def test_freelance_mention_requires_proof():
    cv = make_cv()
    cv.experiences[0].type = "freelance"
    checked, findings = verify(cv, SOURCE, [], apply_removals=True)
    assert checked.experiences[0].type == "emploi" and any("freelance" in f.reason for f in findings)
    checked, _ = verify(cv, SOURCE + "\nBanque Exemple (Freelance)", [], apply_removals=True)
    assert checked.experiences[0].type == "freelance"


def test_linkedin_export_quirks():
    """Export LinkedIn : e-mail coupé en fin de ligne, niveau de langue traduit de l'anglais."""
    source = SOURCE.replace("jean.dupont@example.com", "jean.dupont@exam\nple.com") + "\nFrançais (Native or Bilingual)"
    cv = make_cv(langues=[{"langue": "Français", "niveau": "Langue maternelle ou bilingue"}])
    checked, findings = verify(cv, source, [], apply_removals=True)
    assert checked.contact.email == "jean.dupont@example.com"
    assert not [f for f in findings if "Langue" in f.label]


def _linkedin_pdf(tmp_path: Path):
    """Export LinkedIn fictif : l'entreprise n'est écrite qu'une fois au-dessus de ses deux postes."""
    doc = pymupdf.open()
    page = doc.new_page(width=612, height=792)
    page.insert_text((22, 100), "www.linkedin.com/in/jean-dupont (LinkedIn)", fontsize=11)
    y = 120
    for size, text in [(16, "Expérience"), (12, "ACME Conseil"), (10.5, "3 ans 10 mois"), (11.5, "Chef de projet SI"),
                       (10.5, "octobre 2015 - novembre 2018 (3 ans 2 mois)"), (11.5, "IT Consultant"),
                       (10.5, "février 2015 - septembre 2015 (8 mois)"), (10.5, "Consultant fonctionnel chez Banque Exemple"),
                       (12, "Autre Société"), (11.5, "Stagiaire"), (10.5, "juin 2013 - août 2013 (3 mois)"), (16, "Formation")]:
        page.insert_text((224, y), text, fontsize=size)
        y += 20
    path = tmp_path / "Profile.pdf"
    doc.save(path)
    return load_source(path, with_photos=False)


def test_linkedin_roles_keep_their_company(tmp_path: Path):
    source = _linkedin_pdf(tmp_path)
    assert [(r.company, r.title) for r in source.linkedin_roles] == [
        ("ACME Conseil", "Chef de projet SI"), ("ACME Conseil", "IT Consultant"), ("Autre Société", "Stagiaire")]
    cv = make_cv()
    cv.experiences = [cv.experiences[1].model_copy(update={  # 2e poste attribué au client cité dans sa description
        "poste": "Consultant informatique", "entreprise": "Banque Exemple", "client": None, "type": "emploi",
        "periode_texte": "février 2015 - septembre 2015", "debut": "2015-02", "fin": "2015-09", "realisations": []})]
    checked, findings = verify(cv, source.reference_text, source.links, apply_removals=True, linkedin=source.linkedin_roles)
    exp = checked.experiences[0]
    assert (exp.entreprise, exp.client) == ("ACME Conseil", "Banque Exemple")
    assert any("LinkedIn" in f.reason for f in findings)


def test_non_linkedin_cv_has_no_linkedin_structure(tmp_path: Path):
    assert _pdf(tmp_path, [[(60, EXPERIENCE)]]).linkedin_roles == []


def _pdf(tmp_path: Path, pages: list[list[tuple[float, str]]]):
    """PDF fictif : pour chaque page, des zones de texte (ordonnée, texte multiligne)."""
    doc = pymupdf.open()
    for blocks in pages:
        page = doc.new_page(width=595, height=842)
        for y, text in blocks:
            page.insert_textbox(pymupdf.Rect(40, y, 555, y + 14 * (text.count("\n") + 2)), text, fontsize=11)
    path = tmp_path / "cv.pdf"
    doc.save(path)
    return load_source(path, with_photos=False)


def test_body_text_excludes_header_title_contacts_and_running_footers(tmp_path: Path):
    footer = (815, "Logiclever - Dossier de compétences")
    source = _pdf(tmp_path, [
        [(30, "Mohammed AMMARA\nConsultant confirmé Logiclever\nmohammed@logiclever.com"), (200, EXPERIENCE), footer],
        [(60, "COSUMAR - Raffinerie de Sucre\nStage ingénieur maintenance"), footer],
    ])
    body = source.body_text("Mohammed", "AMMARA", "Consultant confirmé Logiclever")
    assert body is not None and "logiclever" not in body.lower() and "Clem" in body and "COSUMAR" in body
    checked, _ = verify(_ammara(), source.reference_text, source.links, apply_removals=True, body_text=body)
    assert (checked.experiences[0].entreprise, checked.experiences[0].client) == ("Clem’", None)


# --- Contre-vérification -------------------------------------------------------------------


def test_cross_check_removes_atomic_items_but_never_whole_experiences():
    cv = make_cv()
    verdicts = all_confirmed(cv)
    for v in verdicts:
        if v.id in ("E0.P", "E0.R1", "C0"):
            v.statut = "absent"
    findings = apply_verdicts(cv, verdicts, apply_removals=True)
    assert len(cv.experiences) == 2 and cv.experiences[0].poste is None  # seul le poste est vidé
    assert len(cv.experiences[0].realisations) == 1  # E0.R1 absent : retiré
    assert cv.certifications == []  # C0 absent : retiré
    assert {f.action for f in findings} == {"retiré"}


def test_cross_check_inexact_atomic_item_is_removed():
    cv = make_cv()
    apply_verdicts(cv, [Verdict(id="E0.R0", statut="inexact", explication="rattachée à une autre expérience")], apply_removals=True)
    assert cv.experiences[0].realisations == ["Animation des comités de pilotage et suivi du budget de 2,3 M€"]


def test_cross_check_detaches_employer_and_shows_client_as_company():
    cv = _ammara()
    verdicts = [Verdict(id="E0.E", statut="absent", explication="Logiclever n'apparaît que dans le titre"),
                Verdict(id="E0.K", statut="inexact", explication="Clem' est l'entreprise, pas un client")]
    findings = apply_verdicts(cv, verdicts, apply_removals=True)
    assert (cv.experiences[0].entreprise, cv.experiences[0].client) == ("Clem’", None)
    assert any("affiché comme entreprise" in f.reason for f in findings)


def test_cross_check_contested_client_or_dates_are_removed():
    cv = make_cv()
    apply_verdicts(cv, [Verdict(id="E0.K", statut="absent", explication=""), Verdict(id="E0.D", statut="inexact", explication="")], apply_removals=True)
    exp = cv.experiences[0]
    assert exp.entreprise == "ACME Conseil" and exp.client is None
    assert (exp.periode_texte, exp.debut, exp.fin) == (None, None, None)


def test_cross_check_unproven_freelance_mention_is_dropped():
    cv = make_cv()
    cv.experiences[0].type = "freelance"
    apply_verdicts(cv, [Verdict(id="E0.F", statut="absent", explication="")], apply_removals=True)
    assert cv.experiences[0].type == "emploi"


def test_contested_language_level_keeps_the_language():
    cv = make_cv()
    apply_verdicts(cv, [Verdict(id="L0.N", statut="inexact", explication="score listé à part")], apply_removals=True)
    assert [(lang.langue, lang.niveau) for lang in cv.langues] == [("Anglais", None)]


def test_client_only_experience_is_never_presented_as_employer(tmp_path: Path):
    cv = make_cv()
    cv.experiences[0] = cv.experiences[0].model_copy(update={"poste": None, "entreprise": None, "client": "Banque Exemple"})
    claims = " ".join(c.text for c in build_claims(cv) if c.id.startswith("E0."))
    assert "chez « Banque Exemple »" not in claims and "pour le client « Banque Exemple »" in claims
    _, prs = _render(tmp_path, cv, source_pages=1)
    assert "Client : Banque Exemple" in _all_text(prs)


def test_claims_never_anchor_on_the_checked_fact():
    claims = {c.id: c.text for c in build_claims(_ammara())}
    assert "Logiclever" in claims["E0.E"] and "Clem" not in claims["E0.E"]
    assert "Clem" in claims["E0.K"] and "Logiclever" not in claims["E0.K"]
    assert not any(c.startswith(("M", "E0.V")) for c in claims)  # plus de méthodes ni d'environnement


def test_cross_check_accepts_bracketed_ids():
    cv = make_cv()
    verdicts = [Verdict(id="[C0]", statut="absent", explication="introuvable")]
    apply_verdicts(cv, verdicts, apply_removals=True)
    assert cv.certifications == []


def test_cross_check_only_flags_when_reference_is_doubtful():
    cv = make_cv()
    findings = apply_verdicts(cv, [Verdict(id="E0.E", statut="absent", explication="")], apply_removals=False)
    assert cv.experiences[0].entreprise == "ACME Conseil" and findings[0].action == "à vérifier"


def test_nature_contradiction_excludes_from_years_without_removal():
    cv = make_cv()
    cv.experiences[0].type = "emploi"
    findings = apply_verdicts(cv, [Verdict(id="E0.T", statut="inexact", explication="stage")], apply_removals=True)
    assert cv.experiences[0].type == "autre" and len(cv.experiences) == 2
    assert findings and findings[0].action == "à vérifier"


def test_photo_verdict_parsing():
    assert photo_verdict([Verdict(id="[PHOTO]", statut="inexact", explication="badge")]) == (False, "badge")
    assert photo_verdict([Verdict(id="PHOTO", statut="confirme", explication="")]) == (True, "")
    assert photo_verdict([]) == (None, "")


# --- Cache et entrée -----------------------------------------------------------------------------


def _job(stored: dict[str, Any], sha: str = "abc") -> Job:
    job = Job(INPUT_DIR / "cv.pdf", CVReport(source="cv.pdf"), INPUT_DIR / "cv.json")
    job.source = SourceDocument(INPUT_DIR / "cv.pdf", b"", sha, 1, "", "", [])
    job.stored, job.cv = stored, CV.model_validate(stored["cv"])
    return job


def test_json_edited_ignores_schema_changes_and_old_extractions_are_redone():
    data = make_cv().model_dump() | {"methodes": ["Scrum"]}  # champ d'une ancienne version du schéma
    stored: dict[str, Any] = {"_meta": {"sha256": "abc", "empreinte": fingerprint_data(data)}, "cv": data}
    assert not json_edited(stored)
    assert needs_extraction(_job(stored), forcer=False, depuis_json=False)  # consigne antérieure : refaite
    assert not needs_extraction(_job(stored), forcer=False, depuis_json=True)
    stored["_meta"]["version_extraction"] = EXTRACTION_VERSION
    assert not needs_extraction(_job(stored), forcer=False, depuis_json=False)
    edited: dict[str, Any] = {"_meta": {"sha256": "abc", "empreinte": fingerprint_data(data)}, "cv": data | {"titre": "Corrigé à la main"}}
    assert json_edited(edited) and not needs_extraction(_job(edited), forcer=False, depuis_json=False)


def test_input_folder_is_the_default_and_zips_are_unpacked(tmp_path: Path):
    assert default_input() == INPUT_DIR and INPUT_DIR.name == "Input"
    folder = tmp_path / "Input"
    folder.mkdir()
    (folder / "a.pdf").write_bytes(b"%PDF-1.4")
    with zipfile.ZipFile(folder / "drive.zip", "w") as archive:
        archive.writestr("b.pdf", b"%PDF-1.4")
    names = sorted(p.name for p in collect_files([folder], tmp_path / "sortie"))
    assert names == ["a.pdf", "b.pdf"]


# --- Années d'expérience ----------------------------------------------------------------------


@pytest.mark.parametrize(
    "text, expected",
    [
        ("Jan - Déc 2021", ("2021-01", "2021-12")),  # le modèle lisait parfois « 2021 → 2021 » (0 mois)
        ("Apr - Sep 2024", ("2024-04", "2024-09")),
        ("06/2022 – 04/2024", ("2022-06", "2024-04")),
        ("Depuis fev. 2024", ("2024-02", "present")),
        ("septembre 2024 - Aujourd’hui", ("2024-09", "present")),
        ("janvier 2025 - Present (1 an 10 mois)", ("2025-01", "present")),
        ("Septembre 2021 - En cours (CDI)", ("2021-09", "present")),
        ("2023 - Mai 2024", ("2023", "2024-05")),
        ("2016 - 2019", ("2016", "2019")),
        ("Depuis 1 an", (None, None)),  # illisible : on garde la lecture du modèle
    ],
)
def test_period_parsing(text: str, expected: tuple[str | None, str | None]):
    assert contenu.parse_period(text) == expected


def test_period_read_by_code_overrides_model():
    cv = make_cv()
    cv.experiences[0] = cv.experiences[0].model_copy(update={"periode_texte": "Janvier 2020 - Décembre 2021", "debut": "2020", "fin": "2020"})
    checked, _ = verify(cv, SOURCE + "\nJanvier 2020 - Décembre 2021", [], apply_removals=True)
    assert (checked.experiences[0].debut, checked.experiences[0].fin) == ("2020-01", "2021-12")


def test_years_exclude_alternance_and_round_up():
    label, origin, _ = contenu.experience_label(make_cv())
    assert "alternance" in origin
    months = contenu.computed_years(make_cv())[1]
    assert "prouvés" in months
    # mission depuis janvier 2020 : arrondi supérieur
    import datetime as dt

    today = dt.date.today()
    expected = -(-((today.year - 2020) * 12 + today.month) // 12)
    assert label == f"{expected} ans d’expérience"


def test_no_experience_means_no_pill():
    cv = make_cv()
    cv.experiences = [cv.experiences[1]]  # seulement l'alternance
    label, _, _ = contenu.experience_label(cv)
    assert label is None


def test_internship_and_keywords_excluded():
    cv = make_cv()
    cv.experiences[0].type = "emploi"
    cv.experiences[0].poste = "Stage chef de projet"
    assert contenu.computed_years(cv)[0] is None or "stage" in contenu.computed_years(cv)[1]


def test_stated_years_never_exceeded_and_lowered_when_unproven():
    cv = make_cv(annees_experience={"valeur": 2, "citation": "2 ans"})
    assert contenu.experience_label(cv)[0] == "2 ans d’expérience"  # jamais plus que l'annonce
    cv = make_cv(annees_experience={"valeur": 30, "citation": "30 ans"})
    label, _, warning = contenu.experience_label(cv)
    assert label != "30 ans d’expérience" and warning is not None and "ramenées" in warning


def test_founder_role_counts():
    cv = make_cv()
    cv.experiences = [cv.experiences[0].model_copy(update={"type": "creation_entreprise", "debut": "2024-02", "periode_texte": "Depuis février 2024"})]
    years = contenu.computed_years(cv)[0]
    assert years is not None and years >= 1


# --- Photo ----------------------------------------------------------------------------------


def _badge() -> Image.Image:
    img = Image.new("RGB", (400, 400), "white")
    draw = ImageDraw.Draw(img)
    draw.ellipse((10, 10, 390, 390), fill=(255, 93, 36))
    font = ImageFont.truetype(str(FONTS_DIR / "Lexend-SemiBold.ttf"), 46)
    draw.text((200, 200), "CERTIFIED", fill="white", font=font, anchor="mm")
    return img


def test_badge_is_not_a_face():
    assert check_face(_badge()).found is False


@pytest.mark.parametrize(
    "local, llm, edited, keep",
    [
        (True, True, False, True),
        (True, False, False, True),
        (False, True, False, True),
        (False, False, False, False),
        (False, None, False, False),
        (False, False, True, True),
        (None, None, False, True),
        (None, False, False, False),
    ],
)
def test_photo_decision_matrix(local: bool | None, llm: bool | None, edited: bool, keep: bool):
    assert photo_decision(FaceCheck(local, 0.9, "détail"), llm, "", edited)[0] is keep


def test_photo_circle_always_contains_the_face():
    """Photo dans une forme en « œuf » (plus large en bas) : le plus grand cercle de la forme passe sous le visage ;
    le cercle retenu doit contenir le visage, sinon la forme d'origine est conservée (aucun visage coupé)."""
    import numpy as np

    from cv_formatter.pdf_source import best_circle

    ys, xs = np.indices((500, 400))
    egg = ((xs - 200) / (120 + 0.12 * ys)) ** 2 + ((ys - 260) / 240) ** 2 <= 1
    face = (150.0, 60.0, 250.0, 170.0)  # visage en haut de la forme
    found = best_circle(egg, face)
    assert found is not None
    cx, cy, r = found
    assert all(np.hypot(x - cx, y - cy) <= r for x in (face[0], face[2]) for y in (face[1], face[3]))
    assert best_circle(egg, (0.0, 0.0, 399.0, 499.0)) is None  # rectangle impossible à inscrire : on n'en invente pas


def test_pdf_photo_candidate_detected(tmp_path: Path):
    buf = io.BytesIO()
    _badge().save(buf, "PNG")
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    page.insert_image(pymupdf.Rect(40, 40, 160, 160), stream=buf.getvalue())
    for i, line in enumerate(["Jean DUPONT — Chef de projet SI", "Expérience : ACME Conseil, 2020 - 2024", "Compétences : SAP, Power BI"]):
        page.insert_text((40, 200 + 20 * i), line, fontsize=12)  # assez de texte : page non scannée
    pdf = tmp_path / "cv.pdf"
    doc.save(pdf)
    source = load_source(pdf)
    assert len(source.candidates) == 1 and source.text_is_reliable


# --- Typographie / mesure -------------------------------------------------------------------


def test_typography_helpers():
    assert contenu.tame_caps("CHEF DE PROJET SI JUNIOR") == "Chef de projet SI junior"
    assert contenu.tame_caps("PMO IT/IS CONTRAT EN ALTERNANCE") == "PMO IT/IS contrat en alternance"
    assert contenu.tame_caps("Expert MAXIMO/MAS") == "Expert MAXIMO/MAS"
    assert contenu.tame_caps("CONSULTANTE AMOA SI") == "Consultante AMOA SI"
    assert contenu.clean_item("gestion des sprints, etc.") == "Gestion des sprints, etc."
    assert contenu.clean_item("spaCy") == "spaCy"
    assert contenu.clean_item("dbt", capitalize=False) == "dbt"
    assert fix_text("Ing´enieur") == "Ingénieur"
    assert safe_name("CV - Expert MAXIMO/MAS") == "CV - Expert MAXIMO-MAS"


def test_measurement_and_split():
    para = Para([Run("mot " * 40, "regular", 10)], bullet="•", indent=0.4)
    assert line_count(para, 12.44) > 1
    long = Para([Run("mot " * 3000, "regular", 10)])
    pieces = split_long_para(long, 12.44)
    assert len(pieces) > 1 and all(para_height(p, 12.44) <= 12.0 for p in pieces)


def test_measurement_matches_powerpoint_rules():
    # PowerPoint ne fusionne pas les espaces : « Scrum  ·  SAFe » est plus large que « Scrum · SAFe ».
    single, double = "Scrum · SAFe", "Scrum  ·  SAFe"
    middle = (text_width(single, "regular", 10) + text_width(double, "regular", 10)) / 2
    width = middle * (1 + WIDTH_SAFETY_MARGIN) / PT_PER_CM + 2 * TEXT_INSET
    assert line_count(Para([Run(single, "regular", 10)]), width) == 1
    assert line_count(Para([Run(double, "regular", 10)]), width) == 2
    # Insécables avant le point : une ligne coupée finit par « · » au lieu de commencer par lui.
    assert contenu.inline_para(["Scrum", "SAFe", "Jira"])[0].text == "Scrum  ·  SAFe  ·  Jira"
    # PowerPoint arrondit l'espace avant au point entier (1,6 -> 2 ; 8,8 -> 9) : la mesure fait de même.
    paras = [Para([Run("x", "regular", 10)], space_before=2), Para([Run("x", "regular", 10)], space_before=11)]
    assert [p.space_before for p in scale(paras, 0.8)] == [2.0, 9.0]


# --- Rendu PPTX ------------------------------------------------------------------------------


def _shape_texts(shapes: Iterable[BaseShape]) -> Iterator[str]:
    for sh in shapes:
        if isinstance(sh, GroupShape):  # groupe (icône + titre de section)
            yield from _shape_texts(sh.shapes)
        elif isinstance(sh, Shape):
            yield sh.text_frame.text


def _all_text(prs: PresentationFile) -> str:
    return " ".join(text for slide in prs.slides for text in _shape_texts(slide.shapes))


def _render(tmp_path: Path, cv: CV, template: Path = TEMPLATE_PATH, **kwargs: Any) -> tuple[RenderResult, PresentationFile]:
    out = tmp_path / "cv.pptx"
    result = render_cv(cv, None, template, out, experience_label=kwargs.pop("experience_label", None), contact=cv.contact.model_dump(), **kwargs)
    return result, Presentation(str(out))


def test_render_one_page_and_fields(tmp_path: Path):
    cv = make_cv()
    result, prs = _render(tmp_path, cv, experience_label="5 ans d’expérience", source_pages=1)
    text = _all_text(prs)
    assert result.pages == 1 and len(prs.slides) == 1
    assert "{{" not in text and "Jean DUPONT" in text and "SI & outils" in text and "5 ans d’expérience" in text
    assert "ACME Conseil" in text and "Client : Banque Exemple" in text
    rels = [r.target_ref for r in prs.slides[0].part.rels.values() if r.is_external]
    assert "mailto:jean.dupont@example.com" in rels and any("linkedin.com/in/jean-dupont" in r for r in rels)


def test_render_never_shows_methode_nor_environment_and_hides_short_blocks(tmp_path: Path):
    cv = make_cv(outils_si=["Power BI", "Jira"])  # 2 éléments : bloc trop court
    result, prs = _render(tmp_path, cv, source_pages=1)
    text = _all_text(prs)
    assert "MÉTHODE" not in text.upper() and "Environnement" not in text
    assert "SI & outils" not in text and "EXPERTISE" in text
    assert any("SI & outils" in note and "non affiché" in note for note in result.notes)


def test_render_freelance_client_is_labelled(tmp_path: Path):
    cv = make_cv()
    cv.experiences[0] = cv.experiences[0].model_copy(update={"entreprise": None, "client": "Banque Exemple", "type": "freelance"})
    _, prs = _render(tmp_path, cv, source_pages=1)
    text = _all_text(prs)
    assert "Client : Banque Exemple" in text and "Freelance" in text


def test_render_mission_titles_and_clean_linkedin(tmp_path: Path):
    cv = make_cv(contact={"email": None, "telephone": None, "localisation": None,
                          "linkedin": "https://www.linkedin.com/in/jean-dupont?jobid=1234&lipi=urn%3Ali"})
    cv.experiences[0].realisations = ["Mission Banque Exemple :", "Pilotage d'un portefeuille de 200 projets IT"]
    _, prs = _render(tmp_path, cv, source_pages=1)
    text = _all_text(prs)
    assert "linkedin.com/in/jean-dupont" in text and "jobid" not in text
    rels = [r.target_ref for r in prs.slides[0].part.rels.values() if r.is_external]
    assert "https://www.linkedin.com/in/jean-dupont" in rels
    a = "{%s}" % XML_NS["a"]
    title = next(p for p in prs.slides[0].shapes.element.iter(a + "p") if text_of(p) == "Mission Banque Exemple")
    assert title.find(f".//{a}buNone") is not None  # intitulé sans puce


def test_unused_template_placeholder_is_removed(tmp_path: Path):
    prs = Presentation(str(TEMPLATE_PATH))
    box = prs.slides[0].shapes.add_textbox(Emu(0), Emu(0), Emu(100000), Emu(100000))
    box.text_frame.text = "MÉTHODE{{methode}}"  # champ d'un ancien modèle que l'outil ne remplit plus
    template = tmp_path / "ancien_modele.pptx"
    prs.save(str(template))
    result, out = _render(tmp_path, make_cv(), template=template, source_pages=1)
    assert "{{" not in _all_text(out) and "MÉTHODE" not in _all_text(out)
    assert any("{{methode}}" in note for note in result.notes)


def test_render_long_cv_paginates_only_experiences(tmp_path: Path):
    cv = make_cv()
    cv.experiences = [cv.experiences[0].model_copy(update={"realisations": [f"Réalisation numéro {i} " * 6 for i in range(12)]}) for _ in range(8)]
    result, prs = _render(tmp_path, cv, source_pages=4)
    assert result.pages > 1
    for slide in list(prs.slides)[1:]:
        texts = " ".join(_shape_texts(slide.shapes))
        assert "Expériences (suite)" in texts and "Formation" not in texts and "Certifications" not in texts


def _has_text(shape: BaseShape) -> bool:
    return isinstance(shape, Shape) and shape.has_text_frame and bool(shape.text_frame.text.strip())


def _in_slide(group: GroupShape, shape: BaseShape) -> tuple[int, int]:
    """Coin haut gauche (EMU, repère de la diapositive) d'un enfant d'un groupe à l'échelle 1:1."""
    ch_off = group.element.find("p:grpSpPr/a:xfrm/a:chOff", XML_NS)
    assert ch_off is not None
    return group.left + shape.left - int(ch_off.get("x", "0")), group.top + shape.top - int(ch_off.get("y", "0"))


def test_section_headers_are_aligned(tmp_path: Path):
    """Icône calée sur la marge du texte de sa colonne et centrée sur la hauteur de capitale du titre ;
    titre à la même distance de l'icône partout (pages de suite comprises)."""
    cv = make_cv()
    cv.experiences = [cv.experiences[0].model_copy(update={"realisations": [f"Réalisation numéro {i} " * 6 for i in range(12)]}) for _ in range(8)]
    _, prs = _render(tmp_path, cv, source_pages=4)
    margins = {cm(LEFT_X + TEXT_INSET), cm(RIGHT_X + TEXT_INSET)}
    titles: list[str] = []
    for slide in prs.slides:
        for group in (sh for sh in slide.shapes if isinstance(sh, GroupShape)):
            texts = [sh for sh in group.shapes if _has_text(sh)]
            if not texts:  # icône des coordonnées
                continue
            title = texts[0]
            assert isinstance(title, Shape)
            icon = next(sh for sh in group.shapes if sh.shape_id != title.shape_id)
            (ix, iy), (tx, ty) = _in_slide(group, icon), _in_slide(group, title)
            frame = title.text_frame
            size = frame.paragraphs[0].runs[0].font.size
            assert size is not None
            assert ix in margins
            assert abs(iy + icon.height / 2 - (ty + title.height / 2) - cm(cap_center_offset(size.pt))) < 1000  # < 0,03 mm
            assert abs(tx + frame.margin_left - ix - cm(HEADER_TEXT_OFFSET)) < 1000
            assert frame.paragraphs[0].space_after == 0  # l'espace après du modèle faisait remonter le titre
            titles.append(frame.text)
    assert {"Expériences", "Compétences", "Certifications", "Formation", "Expériences (suite)"} <= set(titles)


def test_contact_rows_are_aligned(tmp_path: Path):
    """Texte aligné sur les titres de la colonne droite, icônes sur un même axe et centrées sur leur ligne ;
    avec moins de lignes, le bloc reste centré au même endroit."""

    def rows(cv: CV) -> list[tuple[Shape, BaseShape]]:
        _, prs = _render(tmp_path, cv, source_pages=1)
        shapes = list(prs.slides[0].shapes)
        values = {v for v in cv.contact.model_dump().values() if v}
        boxes = [sh for sh in shapes if isinstance(sh, Shape) and _has_text(sh) and sh.text_frame.text in values]
        icons = [sh for sh in shapes if not _has_text(sh) and sh.top < cm(6.5) and cm(13.8) < sh.left < cm(15)]
        pairs = [(box, min(icons, key=lambda i: abs(i.top + i.height / 2 - box.top - box.height / 2))) for box in boxes]
        return sorted(pairs, key=lambda pair: pair[0].top)

    full = rows(make_cv())
    assert len(full) == 4
    axis = [icon.left + icon.width / 2 for _, icon in full]
    assert max(axis) - min(axis) < 1000
    for box, icon in full:
        size = box.text_frame.paragraphs[0].runs[0].font.size
        assert size is not None
        assert abs(box.left + box.text_frame.margin_left - cm(RIGHT_X + TEXT_INSET + HEADER_TEXT_OFFSET)) < 10
        assert abs(icon.top + icon.height / 2 - (box.top + box.height / 2) - cm(cap_center_offset(size.pt))) < 1000

    def center(pairs: list[tuple[Shape, BaseShape]]) -> float:
        return (pairs[0][0].top + pairs[-1][0].top + pairs[-1][0].height) / 2

    partial = rows(make_cv(contact={"email": "jean.dupont@example.com", "telephone": None, "localisation": "Lyon (69003)", "linkedin": None}))
    assert len(partial) == 2 and abs(center(partial) - center(full)) < 1000


# --- Anonymisation -----------------------------------------------------------------------------


def test_anonymisation_leaks_nothing():
    cv = make_cv()
    cv.experiences[0].realisations.append("Contact : jean.dupont@example.com, 06 12 34 56 78, Jean DUPONT")
    anon = anonymize(cv)
    dump = anon.model_dump_json()
    assert initials(cv) == "J. D."
    for leak in ("jean.dupont@example.com", "06 12 34 56 78", "DUPONT", "linkedin.com/in/jean-dupont"):
        assert leak not in dump
    assert anon.contact.localisation == "Lyon (69003)"


def test_extraction_request_shape():
    """L'appel GPT-6 Luna : PDF joint, store=False, sortie stricte, refus détectés."""
    from cv_formatter.extraction import ExtractionError, check_response

    with pytest.raises(ExtractionError):
        check_response(NS(status="incomplete", incomplete_details=NS(reason="max_output_tokens"), output=[]))
    with pytest.raises(ExtractionError):
        check_response(NS(status="completed", output=[NS(type="message", content=[NS(type="refusal", refusal="non")])]))

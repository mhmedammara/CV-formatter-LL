"""Copie calibrée pour LibreOffice et polices intégrées (sans LibreOffice ni PowerPoint : arithmétique et structure).

Le rendu réel par LibreOffice est contrôlé par tests/test_rendu_linux.py (dans l'image Docker).
"""

from __future__ import annotations

import struct
import zipfile
from pathlib import Path

import pytest
from fontTools.ttLib import TTFont
from pptx import Presentation

from cv_formatter import fonts, libreoffice
from cv_formatter.config import POWERPOINT_BASELINE, TEMPLATE_PATH
from cv_formatter.layout import Para, Run, lines, wrap
from cv_formatter.render_pptx import render_cv
from test_cv_formatter import make_cv

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"


def _rich_cv():
    """CV fictif qui sollicite tous les styles de paragraphe : puces longues, mission, contexte, projets."""
    cv = make_cv()
    first = cv.experiences[0]
    first.contexte = "Programme de transformation du système d'information financier, 40 personnes, budget de 2,3 M€"
    first.realisations = ["Mission Banque Exemple :"] + [
        f"Réalisation {i} : pilotage d'un portefeuille de 200 projets IT avec Power BI et SAP S4/HANA, "
        "animation des comités de pilotage et suivi du budget" for i in range(4)
    ]
    return cv


def _render(tmp_path: Path, name: str = "cv.pptx", **kwargs) -> Path:
    out = tmp_path / name
    cv = kwargs.pop("cv", None) or _rich_cv()
    render_cv(cv, None, TEMPLATE_PATH, out, experience_label="5 ans d’expérience", contact=cv.contact.model_dump(), source_pages=1, **kwargs)
    return out


def _text_shapes(path: Path):
    prs = Presentation(str(path))
    for slide in prs.slides:
        for sp in slide.shapes.element.iter(P + "sp"):
            if libreoffice.has_text(sp):
                yield sp


# --- Césure : chaque caractère est rendu, une seule fois -----------------------------------------------


@pytest.mark.parametrize("text", [
    "Pilotage d'un projet IT stratégique : ateliers utilisateurs, recueil et priorisation des besoins, suivi du projet",
    "SAP  ·  Jira  ·  Power BI  ·  Confluence  ·  ServiceNow  ·  Salesforce",
    "Anticonstitutionnellementanticonstitutionnellementanticonstitutionnellement suite",
    "  espaces  en tête et  doubles  ",
    "",
])
def test_line_breaks_keep_every_character(text: str):
    para = Para([Run(text[:20], "medium", 10), Run(text[20:], "regular", 10)])
    found = lines(para, 6.0)
    assert len(found) == len(wrap(para, 6.0))
    rebuilt, position = "", 0
    for line in found:
        assert line.start >= position and text[position:line.start].strip(" ") == ""  # seuls des espaces sautés
        rebuilt += text[line.start:line.end]
        position = line.end
    assert rebuilt.replace(" ", "") == text.replace(" ", "")


# --- Copie calibrée -------------------------------------------------------------------------------------


def test_every_line_spacing_used_is_calibrated(tmp_path: Path):
    """Un interligne nouveau dans contenu.py doit être mesuré sur PowerPoint (config.POWERPOINT_BASELINE)."""
    spacings = set()
    for sp in _text_shapes(_render(tmp_path)):
        box, _ = libreoffice.read_box(sp)
        spacings |= {round(p.line_spacing, 2) for p in box.paragraphs}
    assert spacings and spacings <= set(POWERPOINT_BASELINE)


def test_spacing_value_converts_exactly():
    for height in range(1, 6000):
        value = libreoffice._spacing_value(height)  # pyright: ignore[reportPrivateUsage]
        assert (value * 254 + 360) // 720 == height  # conversion 1/100 pt -> 1/100 mm de LibreOffice


def test_round_rectangle_text_inset(tmp_path: Path):
    pill = next(sp for sp in _text_shapes(_render(tmp_path)) if "expérience" in "".join(t.text or "" for t in sp.iter(A + "t")))
    box, _ = libreoffice.read_box(pill)
    assert box.shape_inset == pytest.approx(min(box.width, box.height) * 0.5 * 0.29289, abs=1e-6)


def _libreoffice_baselines(sp) -> list[tuple[str, float]]:
    """Lignes de base que LibreOffice donnera à une zone calibrée (règle mesurée, cf. libreoffice._descent)."""
    ext = sp.find(f"{P}spPr/{A}xfrm/{A}ext")
    shape_inset = libreoffice._shape_inset(sp, int(ext.get("cx")) / 360000, int(ext.get("cy")) / 360000)  # pyright: ignore[reportPrivateUsage]
    body_pr = sp.find(f"{P}txBody/{A}bodyPr")
    top = shape_inset * 1000 + int(body_pr.get("tIns")) / 360
    out = []
    for p in sp.iter(A + "p"):
        value = int(p.find(f"{A}pPr/{A}lnSpc/{A}spcPts").get("val"))
        height = (value * 254 + 360) // 720
        runs = [r for r in p.findall(A + "r") if (r.find(A + "t").text or "") != "•"]
        size = max(int(r.find(A + "rPr").get("sz")) / 100 for r in runs) if runs else 10.0
        baseline = top + height - libreoffice._descent(size) - libreoffice.LIBREOFFICE_BASELINE_SHIFT  # pyright: ignore[reportPrivateUsage]
        out.append(("".join(r.find(A + "t").text or "" for r in runs), baseline / 1000))
        top += height
    return out


def test_calibrated_copy_reproduces_powerpoint_lines(tmp_path: Path):
    source = _render(tmp_path)
    copy = tmp_path / "calibre.pptx"
    result = libreoffice.prepare(source, copy)
    assert result.calibrated >= 8 and not result.skipped
    originals = list(_text_shapes(source))
    calibrated = list(_text_shapes(copy))
    assert len(originals) == len(calibrated)
    checked = 0
    for before, after in zip(originals, calibrated):
        box, paragraphs = libreoffice.read_box(before)
        expected = box.powerpoint_lines()
        produced = _libreoffice_baselines(after)
        assert len(produced) == len(expected)
        for item, (text, baseline) in zip(expected, produced):
            para_text = paragraphs[item.paragraph].para.text
            assert text == para_text[item.line.start:item.line.end]  # césure du modèle, mot pour mot
            assert baseline == pytest.approx(item.baseline, abs=0.001)  # au 1/100 mm près
            checked += 1
        assert after.find(f"{P}txBody/{A}bodyPr").get("anchor") == "t"
    assert checked > 40


def test_bullets_become_text_at_the_paragraph_indent(tmp_path: Path):
    copy = tmp_path / "calibre.pptx"
    libreoffice.prepare(_render(tmp_path), copy)
    bullets = [r for sp in _text_shapes(copy) for r in sp.iter(A + "r") if r.find(A + "t").text == "•"]
    assert bullets
    for run in bullets:
        rpr = run.find(A + "rPr")
        assert rpr.find(A + "latin").get("typeface") == "Arial" and int(rpr.get("spc")) > 0
        p = run.getparent()
        assert p.find(f"{A}pPr/{A}buNone") is not None and p.find(f"{A}pPr/{A}buChar") is None
        indent = int(p.find(A + "pPr").get("marL")) / 360000  # cm
        advance = libreoffice.BULLET_ADVANCES[("•", "Arial")] * int(rpr.get("sz")) / 100 + int(rpr.get("spc")) / 100
        assert advance == pytest.approx(indent * 72 / 2.54, abs=0.01)  # texte au retrait, comme derrière une puce


# --- Polices intégrées ----------------------------------------------------------------------------------


def test_eot_wraps_the_whole_font():
    ttf = fonts.font_file("medium").read_bytes()
    data = fonts.eot(ttf)
    eot_size, font_size, version, flags = struct.unpack_from("<IIII", data)
    assert (eot_size, font_size, version, flags) == (len(data), len(ttf), 0x00020001, 0)
    assert struct.unpack_from("<H", data, 34)[0] == 0x504C and data.endswith(ttf)
    assert "Lexend Medium\x00".encode("utf-16-le") in data and struct.unpack_from("<I", data, 28)[0] == TTFont(fonts.font_file("medium"))["OS/2"].usWeightClass


def test_fonts_embedded_like_powerpoint(tmp_path: Path):
    path = _render(tmp_path)
    fonts.embed_in_pptx(path)
    fonts.embed_in_pptx(path)  # relancé : remplacé, pas dupliqué
    with zipfile.ZipFile(path) as z:
        parts = [n for n in z.namelist() if n.startswith("ppt/fonts/")]
        types = z.read("[Content_Types].xml").decode()
    assert len(parts) == 4 and "application/x-fontdata" in types
    prs = Presentation(str(path))
    font_list = prs.element.find(P + "embeddedFontLst")
    children = [el.tag.split("}")[1] for el in prs.element]
    assert children.index("embeddedFontLst") == children.index("notesSz") + 1
    faces = {f.find(P + "font").get("typeface"): sorted(c.tag.split("}")[1] for c in f if c.tag != P + "font") for f in font_list}
    assert faces == {"Lexend": ["bold", "regular"], "Lexend Medium": ["regular"], "Lexend SemiBold": ["bold"]}
    copy = tmp_path / "calibre.pptx"
    libreoffice.prepare(path, copy)  # LibreOffice rend avec les polices installées, pas les copies intégrées
    assert Presentation(str(copy)).element.find(P + "embeddedFontLst") is None

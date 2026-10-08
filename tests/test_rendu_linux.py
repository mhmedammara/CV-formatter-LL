"""Calage du rendu LibreOffice sur PowerPoint : contrôle de bout en bout, dans l'image Docker (Linux).

Les PDF produits par LibreOffice sont comparés, ligne par ligne, aux positions que PowerPoint donne au même
PPTX (modèle de config.POWERPOINT_BASELINE, validé à 0,044 mm près sur les PDF PowerPoint des 14 CV de test) :
même césure, lignes de base et puces à moins de TOLERANCE_MM. À relancer après toute mise à jour de l'image
(LibreOffice, polices) :

    docker run --rm -v "$PWD:/app" -w /app cv-formatter sh -c "pip install --user -q pytest httpx && python -m pytest -q"

Ignoré là où LibreOffice n'est pas installé, et sous Windows (le PDF y est produit par PowerPoint ; forcer avec
CV_FORMATTER_TEST_LIBREOFFICE=1).
"""

from __future__ import annotations

import os
import re
import sys
from dataclasses import dataclass
from pathlib import Path

import pymupdf
import pytest
from pptx import Presentation

from cv_formatter import libreoffice
from cv_formatter.config import EMU_PER_CM, TEMPLATE_PATH
from cv_formatter.export_pdf import PdfExport, _export_with_libreoffice, libreoffice as soffice_path  # pyright: ignore[reportPrivateUsage]
from cv_formatter.layout import PT_PER_CM
from cv_formatter.render_pptx import render_cv
from test_cv_formatter import make_cv

SOFFICE = soffice_path()
pytestmark = pytest.mark.skipif(
    SOFFICE is None or (sys.platform == "win32" and not os.environ.get("CV_FORMATTER_TEST_LIBREOFFICE")),
    reason="rendu LibreOffice contrôlé dans l'image Linux (Docker)",
)
TOLERANCE_MM = 0.05
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"


@dataclass
class Expected:
    page: int
    text: str
    baseline: float  # cm depuis le haut de la page


def _norm(text: str) -> str:
    return re.sub(r"\s+", "", text)


def _origin(sp) -> tuple[float, float]:
    """Position absolue (cm) d'une forme, groupes compris (titres de section : icône + titre)."""
    off = sp.find(f"{P}spPr/{A}xfrm/{A}off")
    x, y = int(off.get("x")), int(off.get("y"))
    parent = sp.getparent()
    while parent is not None and parent.tag == P + "grpSp":
        xfrm = parent.find(f"{P}grpSpPr/{A}xfrm")
        g_off, g_ext = xfrm.find(A + "off"), xfrm.find(A + "ext")
        c_off, c_ext = xfrm.find(A + "chOff"), xfrm.find(A + "chExt")
        sx = int(g_ext.get("cx")) / max(1, int(c_ext.get("cx")))
        sy = int(g_ext.get("cy")) / max(1, int(c_ext.get("cy")))
        x = int(g_off.get("x")) + (x - int(c_off.get("x"))) * sx
        y = int(g_off.get("y")) + (y - int(c_off.get("y"))) * sy
        parent = parent.getparent()
    return x / EMU_PER_CM, y / EMU_PER_CM


def _expected(pptx: Path) -> list[Expected]:
    out: list[Expected] = []
    for page, slide in enumerate(Presentation(str(pptx)).slides):
        for sp in slide.shapes.element.iter(P + "sp"):
            if not libreoffice.has_text(sp):
                continue
            box, paragraphs = libreoffice.read_box(sp)
            _, top = _origin(sp)
            for item in box.powerpoint_lines():
                text = paragraphs[item.paragraph].para.text[item.line.start:item.line.end]
                if text.strip():
                    out.append(Expected(page, _norm(text), top + item.baseline))
    return out


def _rendered(pdf: Path) -> list[tuple[int, str, float, float | None]]:
    """(page, texte, ligne de base du texte, ligne de base de la puce) de chaque ligne du PDF, en cm."""
    out = []
    with pymupdf.open(pdf) as doc:
        for page_number, page in enumerate(doc):
            for block in page.get_text("rawdict")["blocks"]:
                for line in block.get("lines", []):
                    spans = [s for s in line["spans"] if "".join(c["c"] for c in s["chars"]).strip()]
                    if not spans:
                        continue
                    bullet = spans[0] if "".join(c["c"] for c in spans[0]["chars"]).strip() == "•" else None
                    text_spans = spans[1:] if bullet else spans
                    if not text_spans:
                        continue
                    text = "".join(c["c"] for s in text_spans for c in s["chars"])
                    out.append((page_number, _norm(text), text_spans[0]["origin"][1] / PT_PER_CM,
                                bullet["origin"][1] / PT_PER_CM if bullet else None))
    return out


def _cvs():
    riche = make_cv()
    riche.experiences[0].contexte = "Programme de transformation du système d'information financier, 40 personnes"
    riche.experiences[0].realisations = ["Mission Banque Exemple :"] + [
        f"Réalisation {i} : pilotage d'un portefeuille de 200 projets IT avec Power BI et SAP S4/HANA, animation "
        "des comités de pilotage et suivi du budget" for i in range(5)]
    long = make_cv()
    long.experiences = [long.experiences[0].model_copy(update={
        "realisations": [f"Réalisation numéro {i} : " + "conduite du changement et accompagnement des équipes " * 2 for i in range(9)]})
        for _ in range(6)]
    return {"riche": (riche, 1), "long": (long, 3)}


@pytest.mark.parametrize("nom", ["riche", "long"])
def test_rendu_libreoffice_identique_a_powerpoint(tmp_path: Path, nom: str):
    cv, source_pages = _cvs()[nom]
    pptx = tmp_path / f"{nom}.pptx"
    result = render_cv(cv, None, TEMPLATE_PATH, pptx, experience_label="7 ans d’expérience",
                       contact=cv.contact.model_dump(), source_pages=source_pages)
    export = PdfExport()
    errors = _export_with_libreoffice([(pptx, pptx.with_suffix(".pdf"))], SOFFICE or "", export.notes)
    assert not errors and not export.notes
    with pymupdf.open(pptx.with_suffix(".pdf")) as doc:
        assert doc.page_count == result.pages
    expected = _expected(pptx)
    rendered = _rendered(pptx.with_suffix(".pdf"))
    unused = list(rendered)
    worst = 0.0
    for line in expected:
        candidates = [r for r in unused if r[0] == line.page and r[1] == line.text]
        assert candidates, f"ligne absente ou coupée autrement que prévu : {line.text!r}"
        found = min(candidates, key=lambda r: abs(r[2] - line.baseline))
        unused.remove(found)
        delta = abs(found[2] - line.baseline) * 10
        worst = max(worst, delta)
        assert delta <= TOLERANCE_MM, f"{line.text!r} : ligne de base à {delta:.3f} mm de PowerPoint"
        if found[3] is not None:
            assert abs(found[3] - found[2]) * 10 <= TOLERANCE_MM, f"{line.text!r} : puce décalée de la ligne de base"
    assert not unused, f"lignes en trop dans le PDF : {[r[1] for r in unused][:5]}"
    print(f"{nom} : {len(expected)} lignes, écart maximal {worst:.3f} mm")
